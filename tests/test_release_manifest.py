from __future__ import annotations

import csv
import hashlib
import json
import os
import shutil
import subprocess
import sys
import threading
import time
from collections.abc import Iterator
from contextlib import contextmanager, nullcontext
from dataclasses import dataclass, replace
from pathlib import Path

import pyarrow as pa
import pyarrow.parquet as pq
import pytest
from openpyxl import Workbook

from hk_movie_rag import posters, release_data, release_lock, release_manifest
from hk_movie_rag.cli import main
from hk_movie_rag.config import ExpectedCounts, GcpSettings, Settings
from hk_movie_rag.release_lock import ReleaseLockError
from hk_movie_rag.release_manifest import (
    ManifestError,
    build_release_manifest,
    verify_release_manifest,
)

_VALIDATION_CODES = (
    "COLUMN_CONTRACT",
    "RELEASE_ROW_COUNT",
    "RELEASE_ID_UNIQUE",
    "REQUIRED_FIELDS_COMPLETE",
    "DATE_FORMAT_AND_RANGE",
    "RUNTIME_POSITIVE_INTEGER",
    "NO_SENTINELS",
    "CSV_XLSX_EXACT_ALIGNMENT",
    "DATA_TIER_METADATA_ALIGNMENT",
    "SUPPLEMENT_AUDIT_COMPLETE",
    "AUDIT_SOURCE_TRACEABLE",
    "AUDIT_ISSUE_LEDGER_ALIGNMENT",
    "TIER_COUNTS",
    "PILOT_ALIGNMENT",
    "S_TIER_COMPLETE",
    "RELEASE_QUARANTINE_DISJOINT",
    "CANDIDATE_RECONCILIATION",
    "SOURCE_FIDELITY_TITLE_FORMAT",
    "WORKBOOK_ERROR_SCAN",
    "MERGED_LEDGER_ROW_COUNTS",
    "MERGED_LEDGER_ENTITY_ALIGNMENT",
    "MERGED_LEDGER_ENCODING_AND_DATE_PRESERVATION",
)


@dataclass(frozen=True)
class ReleaseFixture:
    root: Path
    settings: Settings
    release_dir: Path
    manifest_path: Path
    source_manifest: Path
    issue_ledger: Path
    artifact_bytes: dict[str, bytes]


@pytest.fixture
def valid_release(tmp_path: Path) -> ReleaseFixture:
    return _make_release(tmp_path)


def test_manifest_verifier_rejects_tampered_artifact(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if verification trusts declared digests instead of current bytes."""
    build_release_manifest(valid_release.settings)
    (valid_release.release_dir / "movies.csv").write_bytes(b"tampered")

    with pytest.raises(ManifestError, match="SHA-256 mismatch"):
        verify_release_manifest(valid_release.manifest_path)


def test_manifest_requires_one_poster_state_per_movie(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if an omitted poster state can still authorize the full movie set."""
    poster_path = valid_release.release_dir / "poster_manifest.parquet"
    rows = pq.read_table(poster_path).to_pylist()
    _write_parquet(poster_path, rows[:-1], pq.read_schema(poster_path))

    with pytest.raises(ManifestError, match="poster rows=1, release movies=2"):
        build_release_manifest(valid_release.settings)
    assert not valid_release.manifest_path.exists()


def test_manifest_is_deterministic_and_uses_only_source_evidence(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if runtime time, host paths, secrets, or unsorted artifacts leak in."""
    first = build_release_manifest(valid_release.settings)
    first_bytes = valid_release.manifest_path.read_bytes()
    second = build_release_manifest(valid_release.settings)
    second_bytes = valid_release.manifest_path.read_bytes()
    payload = json.loads(first_bytes)

    assert second == first
    assert second_bytes == first_bytes
    assert payload["source_generated_at"] == "2026-07-17T04:08:41.762998+00:00"
    assert payload["source_manifest_sha256"] == _sha256(valid_release.source_manifest)
    assert [item["relative_path"] for item in payload["artifacts"]] == sorted(
        item["relative_path"] for item in payload["artifacts"]
    )
    assert payload["expected_counts"] == {
        "audit_rows": 1,
        "pilot_movies": 1,
        "quarantine_movies": 1,
        "release_movies": 2,
        "tier_counts": {"A": 1, "S": 1},
    }
    assert payload["poster_state_counts"] == {
        "content_conflict": 0,
        "machine_passed": 1,
        "missing": 1,
        "placeholder": 0,
    }
    manifest_text = first_bytes.decode("utf-8")
    assert str(valid_release.root) not in manifest_text
    assert "credential" not in manifest_text.casefold()
    assert "access_token" not in manifest_text.casefold()
    assert "\\" not in manifest_text
    assert first.artifact_count == 6


def test_manifest_verifier_enforces_fail_closed_order(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if later checks obscure the first governed blocking failure."""
    build_release_manifest(valid_release.settings)
    original = _load_manifest(valid_release.manifest_path)
    movies_csv = valid_release.release_dir / "movies.csv"

    payload = _deepcopy(original)
    payload["unexpected"] = True
    payload["source_manifest_sha256"] = "0" * 64
    movies_csv.unlink()
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="schema validity"):
        verify_release_manifest(valid_release.manifest_path)

    _restore_release_artifacts(valid_release)
    payload = _deepcopy(original)
    payload["source_manifest_sha256"] = "0" * 64
    movies_csv.unlink()
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="source-manifest identity"):
        verify_release_manifest(valid_release.manifest_path)

    _restore_release_artifacts(valid_release)
    _write_manifest(valid_release.manifest_path, original)
    movies_csv.unlink()
    with pytest.raises(ManifestError, match="authoritative artifact missing"):
        verify_release_manifest(valid_release.manifest_path)

    _restore_release_artifacts(valid_release)
    payload = _deepcopy(original)
    entry = _artifact(payload, "movies.csv")
    entry["size_bytes"] = int(entry["size_bytes"]) + 1
    entry["row_count"] = 999
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="size mismatch"):
        verify_release_manifest(valid_release.manifest_path)

    payload = _deepcopy(original)
    _artifact(payload, "movies.csv")["row_count"] = 999
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="row-count mismatch"):
        verify_release_manifest(valid_release.manifest_path)

    _restore_release_artifacts(valid_release)
    duplicate_rows = _movie_rows()
    duplicate_rows[1]["movie_id"] = duplicate_rows[0]["movie_id"]
    movie_parquet = valid_release.release_dir / "movies.parquet"
    _write_parquet(movie_parquet, duplicate_rows, pq.read_schema(movie_parquet))
    payload = _deepcopy(original)
    _rebind_artifact(payload, valid_release.root, "movies.parquet")
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="unique IDs"):
        verify_release_manifest(valid_release.manifest_path)

    _restore_release_artifacts(valid_release)
    pilot_path = valid_release.release_dir / "pilot_movies.parquet"
    pilot_rows = pq.read_table(pilot_path).to_pylist()
    pilot_rows[0]["movie_id"] = "1999_UNKNOWN_999"
    _write_parquet(pilot_path, pilot_rows, pq.read_schema(pilot_path))
    payload = _deepcopy(original)
    _rebind_artifact(payload, valid_release.root, "pilot_movies.parquet")
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="foreign-key membership"):
        verify_release_manifest(valid_release.manifest_path)

    _restore_release_artifacts(valid_release)
    poster_path = valid_release.release_dir / "poster_manifest.parquet"
    poster_rows = pq.read_table(poster_path).to_pylist()
    poster_rows[0]["quality_status"] = "missing"
    _write_parquet(poster_path, poster_rows, pq.read_schema(poster_path))
    payload = _deepcopy(original)
    _rebind_artifact(payload, valid_release.root, "poster_manifest.parquet")
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="poster-state reconciliation"):
        verify_release_manifest(valid_release.manifest_path)

    _restore_release_artifacts(valid_release)
    _write_issue_ledger(valid_release.issue_ledger, ["1970_GSQ_001"])
    _rebind_source_manifest(valid_release)
    payload = _deepcopy(original)
    payload["source_manifest_sha256"] = _sha256(valid_release.source_manifest)
    _write_manifest(valid_release.manifest_path, payload)
    with pytest.raises(ManifestError, match="Quarantine disjointness"):
        verify_release_manifest(valid_release.manifest_path)


@pytest.mark.parametrize("missing_name", ["movies.csv", "poster_manifest.parquet"])
def test_builder_rejects_missing_authoritative_artifact(
    valid_release: ReleaseFixture, missing_name: str
) -> None:
    """Breaks if a partial canonical Release can publish an authority file."""
    (valid_release.release_dir / missing_name).unlink()

    with pytest.raises(ManifestError, match="authoritative artifact missing"):
        build_release_manifest(valid_release.settings)
    assert not valid_release.manifest_path.exists()


def test_builder_rejects_stale_extra_release_artifact(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if unmanifested Release payloads can coexist with the authority."""
    (valid_release.release_dir / "stale.parquet").write_bytes(b"stale")

    with pytest.raises(ManifestError, match="unexpected authoritative artifact"):
        build_release_manifest(valid_release.settings)
    assert not valid_release.manifest_path.exists()


def test_builder_replaces_manifest_hardlink_without_mutating_target(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if manifest publication follows an existing hard link."""
    protected = valid_release.root / "protected.json"
    protected.write_bytes(b"protected")
    os.link(protected, valid_release.manifest_path)

    build_release_manifest(valid_release.settings)

    assert protected.read_bytes() == b"protected"
    assert not os.path.samefile(protected, valid_release.manifest_path)
    assert not list(valid_release.release_dir.glob(".*.tmp"))


def test_verifier_rejects_linked_artifact_before_reading_outside(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if verification follows an artifact link outside the repository."""
    build_release_manifest(valid_release.settings)
    target = valid_release.root.parent / "outside.csv"
    target.write_bytes((valid_release.release_dir / "movies.csv").read_bytes())
    artifact = valid_release.release_dir / "movies.csv"
    artifact.unlink()
    try:
        artifact.symlink_to(target)
    except OSError as exc:
        pytest.skip(f"symbolic links are unavailable: {exc}")

    with pytest.raises(ManifestError, match="no-follow"):
        verify_release_manifest(valid_release.manifest_path)


def test_builder_rejects_canonical_movies_that_diverge_from_bound_source(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if self-consistent canonical files can replace formal source values."""
    rows = _movie_rows()
    rows[0]["english_title"] = "Fabricated title"
    _write_csv(valid_release.release_dir / "movies.csv", rows)
    movie_parquet = valid_release.release_dir / "movies.parquet"
    _write_parquet(movie_parquet, rows, pq.read_schema(movie_parquet))

    with pytest.raises(ManifestError, match="source reconciliation"):
        build_release_manifest(valid_release.settings)


def test_builder_rejects_rights_unknown_publishable_primary_poster(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if rights-unknown media can be authorized for production serving."""
    path = valid_release.release_dir / "poster_manifest.parquet"
    rows = pq.read_table(path).to_pylist()
    rows[0]["publishable"] = True
    rows[0]["is_primary"] = True
    _write_parquet(path, rows, pq.read_schema(path))

    with pytest.raises(ManifestError, match="poster policy"):
        build_release_manifest(valid_release.settings)


def test_builder_requires_bound_source_validation_22_of_22(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if a hash-bound but failed validation report can authorize Release."""
    valid_release.settings.source_validation.write_text(
        json.dumps(
            {
                "ok": False,
                "checks_passed": 21,
                "checks_total": 22,
                "checks": [{"ok": True}] * 21 + [{"ok": False}],
            }
        ),
        encoding="utf-8",
    )
    _rebind_source_file(valid_release, valid_release.settings.source_validation)

    with pytest.raises(ManifestError, match="source validation.*22/22"):
        build_release_manifest(valid_release.settings)


def test_builder_requires_pilot_ids_to_equal_tier_flags(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if a pilot can move to an unflagged Release movie without detection."""
    path = valid_release.release_dir / "pilot_movies.parquet"
    rows = pq.read_table(path).to_pylist()
    rows[0]["movie_id"] = "1970_YCCS_001"
    _write_parquet(path, rows, pq.read_schema(path))

    with pytest.raises(ManifestError, match="pilot/tier reconciliation"):
        build_release_manifest(valid_release.settings)


def test_builder_rejects_wrong_arrow_types(valid_release: ReleaseFixture) -> None:
    """Breaks if matching column names conceal a changed canonical storage type."""
    path = valid_release.release_dir / "movie_tiers.parquet"
    rows = pq.read_table(path).to_pylist()
    for row in rows:
        row["pilot_movie"] = str(row["pilot_movie"]).lower()
    schema = pa.schema(
        [
            ("movie_id", pa.string()),
            ("tier", pa.string()),
            ("tier_reason", pa.string()),
            ("human_review", pa.string()),
            ("pilot_movie", pa.string()),
        ]
    )
    _write_parquet(path, rows, schema)

    with pytest.raises(ManifestError, match="data schema invalid.*Arrow schema"):
        build_release_manifest(valid_release.settings)


def test_verifier_rejects_unsorted_artifact_array(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if verifier sorting conceals noncanonical manifest bytes."""
    build_release_manifest(valid_release.settings)
    payload = _load_manifest(valid_release.manifest_path)
    artifacts = payload["artifacts"]
    assert isinstance(artifacts, list)
    artifacts.reverse()
    _write_manifest(valid_release.manifest_path, payload)

    with pytest.raises(ManifestError, match="artifacts are not sorted"):
        verify_release_manifest(valid_release.manifest_path)


def test_corrupt_quarantine_workbook_is_governed_error(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if malformed Office ZIP bytes escape the CLI's ManifestError boundary."""
    valid_release.issue_ledger.write_bytes(b"not-an-xlsx")
    _rebind_source_file(valid_release, valid_release.issue_ledger)

    with pytest.raises(ManifestError, match="ledger is unreadable"):
        build_release_manifest(valid_release.settings)


def test_cli_converts_manifest_error_to_status_2(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Breaks if governed manifest failures escape as tracebacks or status 1."""
    valid_release.issue_ledger.write_bytes(b"not-an-xlsx")
    _rebind_source_file(valid_release, valid_release.issue_ledger)
    monkeypatch.chdir(valid_release.root)
    monkeypatch.setattr(sys, "argv", ["hk-movie-rag", "build-manifest"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    captured = capsys.readouterr()
    assert "ledger is unreadable" in captured.err
    assert "Traceback" not in captured.err


def test_builder_rejects_unbound_poster_identity_and_evidence(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if a row can authorize media not bound to its movie and state."""
    path = valid_release.release_dir / "poster_manifest.parquet"
    rows = pq.read_table(path).to_pylist()
    rows[0].update(
        {
            "asset_id": "unbound-asset",
            "source_object_uri": "",
            "original_object_uri": "../outside.jpg",
            "derived_object_uri": "NUL",
            "content_sha256": "",
            "detected_format": "",
            "detected_mime_type": "",
            "detected_extension": "",
            "width": None,
            "height": None,
            "color_mode": "",
            "byte_length": None,
            "rights_status": "cleared",
            "publishable": True,
            "is_primary": True,
        }
    )
    _write_parquet(path, rows, pq.read_schema(path))

    with pytest.raises(ManifestError, match="poster policy"):
        build_release_manifest(valid_release.settings)


def test_builder_verifies_every_declared_source_artifact(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if unconfigured source-manifest evidence can be missing or false."""
    payload = json.loads(valid_release.source_manifest.read_text(encoding="utf-8"))
    payload["artifacts"].append(
        {
            "filename": "contract.pdf",
            "size_bytes": 99,
            "sha256": "f" * 64,
        }
    )
    payload["artifact_count"] = len(payload["artifacts"])
    _write_manifest(valid_release.source_manifest, payload)

    with pytest.raises(ManifestError, match="source-manifest identity.*contract.pdf"):
        build_release_manifest(valid_release.settings)


def test_late_release_file_blocks_manifest_publication(
    valid_release: ReleaseFixture, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Breaks if namespace drift after preflight can coexist with a new authority."""
    original_writer = release_manifest._atomic_replace_bytes

    def insert_late_file(
        root: Path, path: Path, content: bytes, **kwargs: object
    ) -> None:
        (valid_release.release_dir / "late.parquet").write_bytes(b"late")
        original_writer(root, path, content, **kwargs)

    monkeypatch.setattr(release_manifest, "_atomic_replace_bytes", insert_late_file)

    with pytest.raises(ManifestError, match="unexpected authoritative artifact"):
        build_release_manifest(valid_release.settings)
    assert not valid_release.manifest_path.exists()


def test_builder_binds_exact_v12_placeholder_signature(
    valid_release: ReleaseFixture,
) -> None:
    """Breaks if any fabricated repeated hash can replace the known placeholder."""
    path = valid_release.release_dir / "poster_manifest.parquet"
    rows = pq.read_table(path).to_pylist()
    for row in rows:
        movie_id = str(row["movie_id"])
        row.update(
            {
                "source_object_uri": f"{movie_id}.jpg",
                "original_object_uri": "",
                "derived_object_uri": "",
                "quarantine_object_uri": "",
                "content_sha256": "e" * 64,
                "detected_format": "PNG",
                "detected_mime_type": "image/png",
                "detected_extension": ".png",
                "width": 65,
                "height": 91,
                "color_mode": "RGBA",
                "byte_length": 955,
                "quality_status": "placeholder",
            }
        )
    _write_parquet(path, rows, pq.read_schema(path))
    expected = replace(
        valid_release.settings.expected,
        poster_states={
            "machine_passed": 0,
            "content_conflict": 0,
            "placeholder": 2,
            "missing": 0,
        },
    )
    settings = replace(valid_release.settings, expected=expected)

    with pytest.raises(ManifestError, match="placeholder signature"):
        build_release_manifest(settings)


def test_post_install_namespace_failure_removes_new_manifest(
    valid_release: ReleaseFixture, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Breaks if a failed final namespace gate leaves a newly installed authority."""
    original_scan = release_manifest._scan_release_namespace
    calls = 0

    def fail_post_install(
        release_dir: Path,
        *,
        allow_manifest: bool,
        require_manifest: bool = False,
    ) -> None:
        nonlocal calls
        calls += 1
        original_scan(
            release_dir,
            allow_manifest=allow_manifest,
            require_manifest=require_manifest,
        )
        if require_manifest:
            raise ManifestError("unexpected authoritative artifact: late.parquet")

    monkeypatch.setattr(
        release_manifest, "_scan_release_namespace", fail_post_install
    )

    with pytest.raises(ManifestError, match="unexpected authoritative artifact"):
        build_release_manifest(valid_release.settings)
    assert not valid_release.manifest_path.exists()


def test_v12_poster_authority_cannot_be_weakened_by_config() -> None:
    """Breaks if mutable expected counts can disable the frozen v1.2 digest."""
    with pytest.raises(ManifestError, match="frozen poster counts"):
        release_manifest._verify_v12_poster_authority(
            release_version="v1.2",
            release_count=4658,
            expected_counts={
                "machine_passed": 4658,
                "content_conflict": 0,
                "placeholder": 0,
                "missing": 0,
            },
            semantic_sha256="0" * 64,
        )


@pytest.mark.parametrize("mutation", ["delete", "replace", "in_place", "hardlink"])
def test_builder_rebinds_artifacts_immediately_before_publication(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    mutation: str,
) -> None:
    """Breaks if semantic bytes stay authoritative after their paths drift."""
    original_writer = release_manifest._atomic_replace_bytes
    artifact = valid_release.release_dir / "movies.csv"

    def mutate_then_publish(
        root: Path, path: Path, content: bytes, **kwargs: object
    ) -> None:
        _mutate_bound_artifact(artifact, mutation)
        original_writer(root, path, content, **kwargs)

    monkeypatch.setattr(
        release_manifest, "_atomic_replace_bytes", mutate_then_publish
    )

    with pytest.raises(ManifestError, match="artifact.*(?:missing|changed)"):
        build_release_manifest(valid_release.settings)
    assert not valid_release.manifest_path.exists()


@pytest.mark.parametrize("mutation", ["delete", "replace", "in_place", "hardlink"])
def test_verifier_rebinds_artifacts_before_returning_valid(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    mutation: str,
) -> None:
    """Breaks if a late path mutation can race the valid=True boundary."""
    build_release_manifest(valid_release.settings)
    artifact = valid_release.release_dir / "movies.csv"
    original_reader = release_manifest._read_quarantine_ids

    def mutate_after_semantics(content: bytes) -> tuple[set[str], int]:
        result = original_reader(content)
        _mutate_bound_artifact(artifact, mutation)
        return result

    monkeypatch.setattr(
        release_manifest, "_read_quarantine_ids", mutate_after_semantics
    )

    with pytest.raises(ManifestError, match="artifact.*(?:missing|changed)"):
        verify_release_manifest(valid_release.manifest_path)


@pytest.mark.parametrize("mutation", ["delete", "replace", "in_place"])
def test_verifier_rebinds_manifest_before_returning_valid(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    mutation: str,
) -> None:
    """Breaks if valid=True can describe manifest bytes no longer installed."""
    build_release_manifest(valid_release.settings)
    original_reader = release_manifest._read_quarantine_ids

    def mutate_after_semantics(content: bytes) -> tuple[set[str], int]:
        result = original_reader(content)
        _mutate_bound_artifact(valid_release.manifest_path, mutation)
        return result

    monkeypatch.setattr(
        release_manifest, "_read_quarantine_ids", mutate_after_semantics
    )

    with pytest.raises(ManifestError, match="release manifest.*(?:missing|changed)"):
        verify_release_manifest(valid_release.manifest_path)


@pytest.mark.parametrize("prior", [None, b"exact-prior-authority\n"])
def test_post_writer_gate_failure_rolls_back_exact_manifest(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    prior: bytes | None,
) -> None:
    """Breaks if a gate after install falls outside rollback ownership."""
    if prior is not None:
        valid_release.manifest_path.write_bytes(prior)
    original_rebind = release_manifest._rebind_release_artifacts

    def fail_post_install(*args: object, **kwargs: object) -> None:
        if kwargs.get("require_manifest") is True and kwargs.get("lock") is False:
            raise ManifestError("installed release manifest rebind failed")
        original_rebind(*args, **kwargs)

    monkeypatch.setattr(
        release_manifest, "_rebind_release_artifacts", fail_post_install
    )

    with pytest.raises(ManifestError, match="installed release manifest rebind"):
        build_release_manifest(valid_release.settings)
    if prior is None:
        assert not valid_release.manifest_path.exists()
    else:
        assert valid_release.manifest_path.read_bytes() == prior


@pytest.mark.parametrize("operation", ["build", "verify"])
@pytest.mark.parametrize("mutation", ["in_place", "replace"])
def test_authorization_snapshot_blocks_writer_after_first_final_item(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    operation: str,
    mutation: str,
) -> None:
    """Breaks if an already-checked artifact can mutate before authorization."""
    build_release_manifest(valid_release.settings)
    artifact = valid_release.release_dir / "enrichment_provenance.parquet"
    original_bytes = artifact.read_bytes()
    changed_bytes = bytes([original_bytes[0] ^ 1]) + original_bytes[1:]
    mutation_started = threading.Event()
    mutation_done = threading.Event()
    mutation_escaped = False

    def fake_atomic_write(path: Path, content: bytes) -> None:
        del content
        if path.name == artifact.name:
            if mutation == "in_place":
                path.write_bytes(changed_bytes)
            else:
                replacement = path.with_name(f".{path.name}.replacement")
                replacement.write_bytes(changed_bytes)
                os.replace(replacement, path)

    monkeypatch.setattr(release_data, "_atomic_replace_bytes", fake_atomic_write)

    def mutate_from_governed_writer() -> None:
        mutation_started.set()
        release_data._write_outputs(
            valid_release.release_dir,
            movies=[],
            tiers=[],
            pilots=[],
            provenance=[],
        )
        mutation_done.set()

    writer = threading.Thread(target=mutate_from_governed_writer, daemon=True)
    attempt_started = False

    def begin_attempt() -> None:
        nonlocal attempt_started, mutation_escaped
        if attempt_started:
            return
        attempt_started = True
        writer.start()
        assert mutation_started.wait(timeout=2)
        mutation_escaped = mutation_done.wait(timeout=0.25)

    original_read = release_manifest._read_bound_file
    artifact_reads = 0
    trigger_read = 13 if operation == "build" else 7

    def read_then_attempt(
        root: Path, path: Path, label: str
    ) -> tuple[bytes, object]:
        nonlocal artifact_reads
        result = original_read(root, path, label)
        if label.startswith("artifact "):
            artifact_reads += 1
            if artifact_reads == trigger_read:
                begin_attempt()
        return result

    monkeypatch.setattr(release_manifest, "_read_bound_file", read_then_attempt)
    original_held_verify = getattr(release_manifest, "_verify_held_file", None)
    held_verifications = 0
    trigger_held_verification = 8 if operation == "build" else 1

    def held_then_attempt(*args: object, **kwargs: object) -> None:
        nonlocal held_verifications
        assert original_held_verify is not None
        original_held_verify(*args, **kwargs)
        held_verifications += 1
        if held_verifications == trigger_held_verification:
            begin_attempt()

    monkeypatch.setattr(
        release_manifest, "_verify_held_file", held_then_attempt, raising=False
    )

    if operation == "build":
        build_release_manifest(valid_release.settings)
    else:
        assert verify_release_manifest(valid_release.manifest_path).valid is True

    assert attempt_started
    assert mutation_escaped is False
    writer.join(timeout=5)
    assert mutation_done.is_set()
    with pytest.raises(ManifestError, match="SHA-256 mismatch"):
        verify_release_manifest(valid_release.manifest_path)


def test_release_lock_is_cross_process(valid_release: ReleaseFixture) -> None:
    """Breaks if the release guard coordinates only threads or one process."""
    ready = valid_release.root / "child-ready"
    completed = valid_release.root / "child-completed"
    guard_factory = getattr(release_manifest, "_release_guard", None)
    guard = (
        guard_factory(valid_release.root, "v0.1", exclusive=True)
        if guard_factory is not None
        else nullcontext()
    )
    script = """
import contextlib
import sys
from pathlib import Path

try:
    from hk_movie_rag.release_lock import release_guard
except ImportError:
    def release_guard(*args, **kwargs):
        return contextlib.nullcontext()

root = Path(sys.argv[1])
ready = Path(sys.argv[2])
completed = Path(sys.argv[3])
ready.write_text("ready", encoding="utf-8")
with release_guard(root, "v0.1", exclusive=True):
    completed.write_text("completed", encoding="utf-8")
"""
    process: subprocess.Popen[str] | None = None
    with guard:
        process = subprocess.Popen(
            [sys.executable, "-c", script, str(valid_release.root), str(ready), str(completed)],
            cwd=Path(__file__).resolve().parents[1],
            text=True,
        )
        deadline = time.monotonic() + 5
        while not ready.exists() and time.monotonic() < deadline:
            time.sleep(0.01)
        assert ready.exists()
        time.sleep(0.25)
        child_was_blocked = not completed.exists()
    assert process is not None
    assert process.wait(timeout=5) == 0
    assert child_was_blocked
    assert completed.exists()


def test_poster_manifest_writer_joins_release_publication_lock(
    valid_release: ReleaseFixture, monkeypatch: pytest.MonkeyPatch
) -> None:
    calls: list[tuple[Path, str, bool]] = []

    @contextmanager
    def recording_guard(
        root: Path, release_version: str, *, exclusive: bool
    ) -> Iterator[None]:
        calls.append((root, release_version, exclusive))
        yield

    expected = object()
    monkeypatch.setattr(posters, "release_guard", recording_guard)
    monkeypatch.setattr(
        posters,
        "_materialize_posters_locked",
        lambda scan, root: expected,
    )
    scan = posters.PosterScanResult(
        release_version="v0.1",
        rows=(),
        present_keys=0,
        orphan_keys=frozenset(),
        detected_formats={},
    )

    assert posters.materialize_posters(scan, valid_release.root) is expected
    assert calls == [(valid_release.root.resolve(), "v0.1", True)]


@pytest.mark.parametrize("prior", [None, b"exact-prior-authority\n"])
def test_rollback_never_clobbers_competing_manifest(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    prior: bytes | None,
) -> None:
    """Breaks if rollback owns a pathname instead of its failed candidate inode."""
    if prior is not None:
        valid_release.manifest_path.write_bytes(prior)
    competing = b"competing-successful-publisher\n"
    original_rebind = release_manifest._rebind_release_artifacts
    post_install_gate_failed = False
    candidate_was_installed = False

    def fail_current_post_install(*args: object, **kwargs: object) -> None:
        nonlocal candidate_was_installed, post_install_gate_failed
        if kwargs.get("require_manifest") is True and kwargs.get("lock") is False:
            candidate_was_installed = valid_release.manifest_path.is_file()
            post_install_gate_failed = True
            raise ManifestError("installed candidate failed its final gate")
        original_rebind(*args, **kwargs)

    monkeypatch.setattr(
        release_manifest, "_rebind_release_artifacts", fail_current_post_install
    )
    original_rollback = release_manifest._rollback_release_manifest
    candidate_owned_rollback_ran = False

    def install_competitor_then_rollback(*args: object, **kwargs: object) -> None:
        nonlocal candidate_owned_rollback_ran
        assert kwargs["failed_binding"] is not None
        assert kwargs["failed_content"] is not None
        replacement = valid_release.manifest_path.with_name(".competing-manifest")
        replacement.write_bytes(competing)
        os.replace(replacement, valid_release.manifest_path)
        original_rollback(*args, **kwargs)
        candidate_owned_rollback_ran = True

    monkeypatch.setattr(
        release_manifest,
        "_rollback_release_manifest",
        install_competitor_then_rollback,
    )

    with pytest.raises(ManifestError):
        build_release_manifest(valid_release.settings)
    assert candidate_was_installed
    assert post_install_gate_failed
    assert candidate_owned_rollback_ran
    assert valid_release.manifest_path.read_bytes() == competing


def test_v12_rejects_coordinated_source_and_config_drift() -> None:
    """Breaks if coordinated mutable inputs can redefine the v1.2 authority."""
    with pytest.raises(ManifestError, match="frozen source manifest"):
        release_manifest._verify_v12_frozen_authority(
            release_version="v1.2",
            source_manifest_sha256="f" * 64,
            source_generated_at="2030-01-01T00:00:00+00:00",
            expected_counts={
                "release_movies": 1,
                "quarantine_movies": 0,
                "tier_counts": {"S": 1},
                "pilot_movies": 0,
                "audit_rows": 0,
            },
            canonical_digests={},
        )


def test_v12_rejects_nonposter_semantic_drift() -> None:
    """Breaks if coordinated canonical bytes can replace frozen movie semantics."""
    canonical_digests = {
        "enrichment_provenance.parquet": "04e5b5712c7f8ee04e7696ee08b5d2821a6108759948c134b3ef237d57506ba0",
        "movie_tiers.parquet": "0070331c8ed428f4702f698ba94a653c8388183803ee632f37a5214deeb755c9",
        "movies.csv": "a0ce946130ff9e53042004ff9ca68a90b399fc81c949d5896df799a5c40dd567",
        "movies.parquet": "a0ce946130ff9e53042004ff9ca68a90b399fc81c949d5896df799a5c40dd567",
        "pilot_movies.parquet": "0" * 64,
    }
    with pytest.raises(ManifestError, match="frozen canonical semantics"):
        release_manifest._verify_v12_frozen_authority(
            release_version="v1.2",
            source_manifest_sha256="cbd529e126e8b72dccac302d9746fcfa63318caecbcda24548c05e282e13f974",
            source_generated_at="2026-07-17T04:08:41.762998+00:00",
            expected_counts={
                "release_movies": 4658,
                "quarantine_movies": 1264,
                "tier_counts": {"S": 50, "A": 313, "B": 4295},
                "pilot_movies": 24,
                "audit_rows": 1008,
            },
            canonical_digests=canonical_digests,
        )


@pytest.mark.parametrize(
    "mutation",
    ["duplicate", "wrong_identity", "wrong_order", "wrong_version", "wrong_severity", "extra_field"],
)
def test_source_validation_requires_exact_22_check_identities(mutation: str) -> None:
    """Breaks if 22 arbitrary truthy dictionaries can impersonate validation."""
    report = _validation_report("v1.2")
    checks = report["checks"]
    assert isinstance(checks, list)
    if mutation == "duplicate":
        checks[1]["code"] = checks[0]["code"]
    elif mutation == "wrong_identity":
        checks[0]["code"] = "DUMMY"
    elif mutation == "wrong_order":
        checks[0], checks[1] = checks[1], checks[0]
    elif mutation == "wrong_version":
        report["version"] = "v9.9"
    elif mutation == "wrong_severity":
        checks[0]["severity"] = "notice"
    else:
        checks[0]["dummy"] = True

    with pytest.raises(ManifestError, match="source validation.*22 identities"):
        release_manifest._validate_source_validation(
            json.dumps(report).encode("utf-8"), "v1.2"
        )


@pytest.mark.parametrize("command", ["build-manifest", "verify-release"])
def test_cli_converts_malformed_yaml_to_status_2(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    command: str,
) -> None:
    """Breaks if malformed YAML escapes the governed argparse boundary."""
    if command == "verify-release":
        build_release_manifest(valid_release.settings)
    (valid_release.root / "config" / "local.yaml").write_text(
        "expected: [\n", encoding="utf-8"
    )
    argv = ["hk-movie-rag", command]
    if command == "verify-release":
        argv.append(str(valid_release.manifest_path))
    monkeypatch.chdir(valid_release.root)
    monkeypatch.setattr(sys, "argv", argv)

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    captured = capsys.readouterr()
    assert "configuration YAML is invalid" in captured.err
    assert "Traceback" not in captured.err


def test_verify_release_cli_accepts_an_inventory_root_detached_after_cleanup(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    """Breaks if Release verification still requires archived raw inventory roots."""
    build_release_manifest(valid_release.settings)
    local_config = valid_release.root / "config" / "local.yaml"
    local_config.write_text(
        local_config.read_text(encoding="utf-8").replace(
            "inventory_roots: [FinalDelivery_2026-07-16]",
            "inventory_roots: [FinalDelivery_2026-07-16, detached-posters]",
        ),
        encoding="utf-8",
    )
    monkeypatch.chdir(valid_release.root)
    monkeypatch.setattr(
        sys,
        "argv",
        ["hk-movie-rag", "verify-release", str(valid_release.manifest_path)],
    )

    main()

    assert json.loads(capsys.readouterr().out) == {
        "artifact_count": 6,
        "movie_count": 2,
        "release_version": "v0.1",
        "valid": True,
    }


@pytest.mark.parametrize("command", ["build-manifest", "verify-release"])
@pytest.mark.parametrize("phase", ["acquire", "release"])
def test_cli_converts_release_lock_failures_to_status_2(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    command: str,
    phase: str,
) -> None:
    if command == "verify-release":
        build_release_manifest(valid_release.settings)

    @contextmanager
    def failing_guard(*args: object, **kwargs: object) -> Iterator[None]:
        del args, kwargs
        if phase == "acquire":
            raise ReleaseLockError("synthetic acquire failure")
        yield
        raise ReleaseLockError("synthetic release failure")

    monkeypatch.setattr(release_manifest, "_release_guard", failing_guard)
    argv = ["hk-movie-rag", command]
    if command == "verify-release":
        argv.append(str(valid_release.manifest_path))
    monkeypatch.chdir(valid_release.root)
    monkeypatch.setattr(sys, "argv", argv)

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    captured = capsys.readouterr()
    assert f"synthetic {phase} failure" in captured.err
    assert "Traceback" not in captured.err


class _CloseFailingLockFile:
    def __init__(self, inner: object) -> None:
        self._inner = inner

    def fileno(self) -> int:
        return self._inner.fileno()  # type: ignore[attr-defined,no-any-return]

    def seek(self, offset: int) -> int:
        return self._inner.seek(offset)  # type: ignore[attr-defined,no-any-return]

    def close(self) -> None:
        self._inner.close()  # type: ignore[attr-defined]
        raise OSError("synthetic raw close failure")


@pytest.mark.parametrize("command", ["build-manifest", "verify-release"])
@pytest.mark.parametrize("phase", ["open", "wrap", "close"])
def test_cli_normalizes_raw_lock_os_errors_to_status_2(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    command: str,
    phase: str,
) -> None:
    """Breaks if a lower-level lock OSError escapes the governed CLI boundary."""
    if command == "verify-release":
        build_release_manifest(valid_release.settings)

    if phase == "open":
        def fail_open(*args: object, **kwargs: object) -> object:
            del args, kwargs
            raise OSError("synthetic raw open failure")

        monkeypatch.setattr(release_lock, "_open_lock_file", fail_open)
    else:
        original_fdopen = release_lock.os.fdopen
        injected = False

        def fail_or_wrap_fdopen(*args: object, **kwargs: object) -> object:
            nonlocal injected
            if injected:
                return original_fdopen(*args, **kwargs)
            injected = True
            if phase == "wrap":
                raise OSError("synthetic raw wrap failure")
            return _CloseFailingLockFile(original_fdopen(*args, **kwargs))

        monkeypatch.setattr(release_lock.os, "fdopen", fail_or_wrap_fdopen)

    argv = ["hk-movie-rag", command]
    if command == "verify-release":
        argv.append(str(valid_release.manifest_path))
    monkeypatch.chdir(valid_release.root)
    monkeypatch.setattr(sys, "argv", argv)

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    captured = capsys.readouterr()
    assert f"synthetic raw {phase} failure" in captured.err
    assert "Traceback" not in captured.err


@pytest.mark.parametrize("prior", [None, b"exact-prior-authority\n"])
def test_build_cli_rolls_back_candidate_when_guard_exit_close_fails(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    prior: bytes | None,
) -> None:
    """Breaks if guard-exit failure leaves its installed manifest authoritative."""
    if prior is not None:
        valid_release.manifest_path.write_bytes(prior)
    original_fdopen = release_lock.os.fdopen
    injected = False

    def close_failing_fdopen(*args: object, **kwargs: object) -> object:
        nonlocal injected
        if injected:
            return original_fdopen(*args, **kwargs)
        injected = True
        return _CloseFailingLockFile(original_fdopen(*args, **kwargs))

    monkeypatch.setattr(release_lock.os, "fdopen", close_failing_fdopen)
    monkeypatch.chdir(valid_release.root)
    monkeypatch.setattr(sys, "argv", ["hk-movie-rag", "build-manifest"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    captured = capsys.readouterr()
    assert "synthetic raw close failure" in captured.err
    assert "Traceback" not in captured.err
    if prior is None:
        assert not valid_release.manifest_path.exists()
    else:
        assert valid_release.manifest_path.read_bytes() == prior


@pytest.mark.parametrize("prior", [None, b"exact-prior-authority\n"])
@pytest.mark.parametrize("schedule", ["rollback_first", "competitor_first"])
def test_guard_exit_rollback_serializes_with_real_competing_publisher(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
    prior: bytes | None,
    schedule: str,
) -> None:
    """Breaks if guard-exit rollback can race a governed publisher."""
    build_release_manifest(valid_release.settings)
    failed_candidate = valid_release.manifest_path.read_bytes()
    competing_authority = failed_candidate.rstrip(b"\n") + b" \n"
    if prior is None:
        valid_release.manifest_path.unlink()
    else:
        valid_release.manifest_path.write_bytes(prior)

    publisher_thread = threading.current_thread()
    allow_competitor_lock = threading.Event()
    competitor_waiting = threading.Event()
    competitor_kernel_wait_started = threading.Event()
    competitor_entered = threading.Event()
    competitor_finished = threading.Event()
    rollback_reacquired = threading.Event()
    competitor_errors: list[BaseException] = []
    competitor_observed: list[bytes | None] = []
    competitor_bindings: list[tuple[int, int]] = []
    failed_candidate_bindings: list[tuple[int, int]] = []
    rollback_was_guarded: list[bool] = []
    completion_order: list[str] = []
    real_guard = release_lock.release_guard
    original_lock_file = release_lock._lock_file
    rollback_guard_active = False
    publication_guard_calls = 0

    def scheduled_lock_file(lock_file: object, *, exclusive: bool) -> None:
        if threading.current_thread() is competitor:
            competitor_waiting.set()
            if not allow_competitor_lock.wait(timeout=5):
                raise AssertionError("competitor lock schedule was never released")
            competitor_kernel_wait_started.set()
        original_lock_file(lock_file, exclusive=exclusive)  # type: ignore[arg-type]

    def competing_publisher() -> None:
        try:
            with real_guard(valid_release.root, "v0.1", exclusive=True):
                competitor_entered.set()
                competitor_observed.append(
                    valid_release.manifest_path.read_bytes()
                    if valid_release.manifest_path.exists()
                    else None
                )
                release_manifest._atomic_replace_bytes(
                    valid_release.root,
                    valid_release.manifest_path,
                    competing_authority,
                )
                result = release_manifest._verify_release_manifest_locked(
                    valid_release.manifest_path,
                    valid_release.root,
                    "v0.1",
                )
                assert result.valid is True
                installed = valid_release.manifest_path.stat()
                competitor_bindings.append((installed.st_dev, installed.st_ino))
                completion_order.append("competitor")
        except (AssertionError, ManifestError, OSError, posters.PosterError) as exc:
            competitor_errors.append(exc)
        finally:
            competitor_finished.set()

    competitor = threading.Thread(
        target=competing_publisher,
        name=f"governed-competitor-{schedule}",
        daemon=True,
    )

    @contextmanager
    def scheduled_publication_guard(
        root: Path, release_version: str, *, exclusive: bool
    ) -> Iterator[None]:
        nonlocal publication_guard_calls, rollback_guard_active
        if threading.current_thread() is not publisher_thread:
            with real_guard(root, release_version, exclusive=exclusive):
                yield
            return

        publication_guard_calls += 1
        call = publication_guard_calls
        if call == 2 and schedule == "competitor_first":
            assert competitor_finished.wait(timeout=5)
        with real_guard(root, release_version, exclusive=exclusive):
            if call == 2:
                rollback_guard_active = True
                rollback_reacquired.set()
            try:
                yield
            finally:
                should_release_competitor = (
                    call == 1 and schedule == "competitor_first"
                ) or (call == 2 and schedule == "rollback_first")
                if should_release_competitor:
                    allow_competitor_lock.set()
                    assert competitor_kernel_wait_started.wait(timeout=5)
                    assert not competitor_entered.is_set()
                if call == 2:
                    rollback_guard_active = False

    original_rebind = release_manifest._rebind_release_artifacts
    competitor_started = False

    def start_competitor_after_install(*args: object, **kwargs: object) -> None:
        nonlocal competitor_started
        original_rebind(*args, **kwargs)
        if (
            threading.current_thread() is publisher_thread
            and kwargs.get("require_manifest") is True
            and kwargs.get("lock") is False
            and not competitor_started
        ):
            competitor_started = True
            installed = valid_release.manifest_path.stat()
            failed_candidate_bindings.append((installed.st_dev, installed.st_ino))
            competitor.start()
            assert competitor_waiting.wait(timeout=5)
            assert not competitor_entered.is_set()

    original_rollback = release_manifest._rollback_release_manifest

    def observe_guarded_rollback(*args: object, **kwargs: object) -> None:
        guarded = rollback_guard_active
        rollback_was_guarded.append(guarded)
        if schedule == "competitor_first" and not guarded:
            assert competitor_finished.wait(timeout=5)
        try:
            original_rollback(*args, **kwargs)
            completion_order.append("rollback")
        finally:
            if schedule == "rollback_first" and not guarded:
                allow_competitor_lock.set()

    original_fdopen = release_lock.os.fdopen
    close_failure_injected = False

    def close_failing_fdopen(*args: object, **kwargs: object) -> object:
        nonlocal close_failure_injected
        opened = original_fdopen(*args, **kwargs)
        if close_failure_injected:
            return opened
        close_failure_injected = True
        return _CloseFailingLockFile(opened)

    monkeypatch.setattr(release_lock, "_lock_file", scheduled_lock_file)
    monkeypatch.setattr(release_manifest, "_release_guard", scheduled_publication_guard)
    monkeypatch.setattr(
        release_manifest, "_rebind_release_artifacts", start_competitor_after_install
    )
    monkeypatch.setattr(
        release_manifest, "_rollback_release_manifest", observe_guarded_rollback
    )
    monkeypatch.setattr(release_lock.os, "fdopen", close_failing_fdopen)
    monkeypatch.chdir(valid_release.root)
    monkeypatch.setattr(sys, "argv", ["hk-movie-rag", "build-manifest"])

    try:
        with pytest.raises(SystemExit) as exc_info:
            main()
    finally:
        allow_competitor_lock.set()
        if competitor_started:
            competitor_finished.wait(timeout=5)
            competitor.join(timeout=1)

    assert exc_info.value.code == 2
    captured = capsys.readouterr()
    assert "synthetic raw close failure" in captured.err
    assert "Traceback" not in captured.err
    assert publication_guard_calls == 2
    assert rollback_reacquired.is_set()
    assert rollback_was_guarded == [True]
    assert competitor_errors == []
    assert competitor_finished.is_set()
    assert len(failed_candidate_bindings) == 1
    assert len(competitor_bindings) == 1
    assert failed_candidate_bindings != competitor_bindings
    if schedule == "rollback_first":
        assert completion_order == ["rollback", "competitor"]
        assert competitor_observed == [prior]
    else:
        assert completion_order == ["competitor", "rollback"]
        assert competitor_observed == [failed_candidate]
    assert valid_release.manifest_path.read_bytes() == competing_authority
    final = valid_release.manifest_path.stat()
    assert (final.st_dev, final.st_ino) == competitor_bindings[0]
    assert valid_release.manifest_path.read_bytes() != failed_candidate
    assert verify_release_manifest(valid_release.manifest_path).valid is True


def test_guard_exit_rollback_reacquire_failure_never_touches_new_authority(
    valid_release: ReleaseFixture,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Breaks if failed rollback reacquisition still mutates the manifest name."""
    build_release_manifest(valid_release.settings)
    failed_candidate = valid_release.manifest_path.read_bytes()
    competing_authority = failed_candidate.rstrip(b"\n") + b" \n"
    valid_release.manifest_path.unlink()
    guard_calls = 0
    rollback_called = False

    @contextmanager
    def failing_reacquire_guard(*args: object, **kwargs: object) -> Iterator[None]:
        nonlocal guard_calls
        del args, kwargs
        guard_calls += 1
        if guard_calls == 1:
            yield
            release_manifest._atomic_replace_bytes(
                valid_release.root,
                valid_release.manifest_path,
                competing_authority,
            )
            raise ReleaseLockError("synthetic guard exit failure")
        raise ReleaseLockError("synthetic rollback reacquire failure")

    def forbidden_rollback(*args: object, **kwargs: object) -> None:
        nonlocal rollback_called
        del args, kwargs
        rollback_called = True

    monkeypatch.setattr(release_manifest, "_release_guard", failing_reacquire_guard)
    monkeypatch.setattr(
        release_manifest, "_rollback_release_manifest", forbidden_rollback
    )

    with pytest.raises(ReleaseLockError, match="rollback reacquire"):
        build_release_manifest(valid_release.settings)

    assert guard_calls == 2
    assert rollback_called is False
    assert valid_release.manifest_path.read_bytes() == competing_authority


def _make_release(root: Path) -> ReleaseFixture:
    release_dir = root / "data" / "release" / "v0.1"
    schema_dir = root / "schemas"
    source_dir = root / "FinalDelivery_2026-07-16"
    config_dir = root / "config"
    for path in (release_dir, schema_dir, source_dir, config_dir):
        path.mkdir(parents=True, exist_ok=True)
    actual_schema_dir = Path(__file__).resolve().parents[1] / "schemas"
    for name in (
        "movie.schema.json",
        "media_asset.schema.json",
        "release_manifest.schema.json",
    ):
        shutil.copy2(actual_schema_dir / name, schema_dir / name)

    _write_csv(release_dir / "movies.csv", _movie_rows())
    _write_parquet(
        release_dir / "movies.parquet",
        _movie_rows(),
        pa.schema([(name, pa.string()) for name in _movie_rows()[0]]),
    )
    _write_parquet(
        release_dir / "movie_tiers.parquet",
        [
            {
                "movie_id": "1970_GSQ_001",
                "tier": "S",
                "tier_reason": "reason",
                "human_review": "required",
                "pilot_movie": True,
            },
            {
                "movie_id": "1970_YCCS_001",
                "tier": "A",
                "tier_reason": "reason",
                "human_review": "not_required",
                "pilot_movie": False,
            },
        ],
        pa.schema(
            [
                ("movie_id", pa.string()),
                ("tier", pa.string()),
                ("tier_reason", pa.string()),
                ("human_review", pa.string()),
                ("pilot_movie", pa.bool_()),
            ]
        ),
    )
    _write_parquet(
        release_dir / "pilot_movies.parquet",
        [
            {
                "movie_id": "1970_GSQ_001",
                "handbook_genre": "drama",
                "evidence_note": "note",
                "evidence_url": "https://example.com/evidence",
                "source_record_id": "record-1",
            }
        ],
        _strings_schema(
            "movie_id",
            "handbook_genre",
            "evidence_note",
            "evidence_url",
            "source_record_id",
        ),
    )
    _write_parquet(
        release_dir / "enrichment_provenance.parquet",
        [
            {
                "movie_id": "1970_YCCS_001",
                "chinese_title": "玉女嬉春",
                "release_date": "1970-01-02",
                "field_name": "english_title",
                "original_value": "",
                "filled_value": "Spring",
                "source_label": "source",
                "source_url": "https://example.com/source",
                "source_record_id": "record-2",
                "match_method": "exact",
                "confidence": "1",
                "matched_imdb_id": "tt0000001",
                "retrieved_at": "2026-07-17T00:00:00.000000",
                "notes": "",
            }
        ],
        _strings_schema(
            "movie_id",
            "chinese_title",
            "release_date",
            "field_name",
            "original_value",
            "filled_value",
            "source_label",
            "source_url",
            "source_record_id",
            "match_method",
            "confidence",
            "matched_imdb_id",
            "retrieved_at",
            "notes",
        ),
    )
    _write_parquet(
        release_dir / "poster_manifest.parquet",
        _poster_rows(),
        _poster_schema(),
    )

    issue_ledger = source_dir / "issue-ledger.xlsx"
    _write_issue_ledger(issue_ledger, ["1980_QUARANTINE_001"])
    placeholders = {
        "movies.csv": source_dir / "source-movies.csv",
        "movies.xlsx": source_dir / "source-movies.xlsx",
        "tiers.xlsx": source_dir / "source-tiers.xlsx",
        "validation.json": source_dir / "validation.json",
    }
    _write_source_movie_csv(placeholders["movies.csv"])
    _write_source_movies_workbook(placeholders["movies.xlsx"])
    _write_source_tiers_workbook(placeholders["tiers.xlsx"])
    placeholders["validation.json"].write_text(
        json.dumps(_validation_report("v0.1"), sort_keys=True),
        encoding="utf-8",
    )
    source_manifest = source_dir / "manifest_v1.2.json"
    _write_source_manifest(source_manifest, issue_ledger, placeholders)

    (config_dir / "local.yaml").write_text(
        """release_version: v0.1
source_manifest: FinalDelivery_2026-07-16/manifest_v1.2.json
source_validation: FinalDelivery_2026-07-16/validation.json
movies_csv: FinalDelivery_2026-07-16/source-movies.csv
movies_xlsx: FinalDelivery_2026-07-16/source-movies.xlsx
tiers_xlsx: FinalDelivery_2026-07-16/source-tiers.xlsx
issue_ledger_xlsx: FinalDelivery_2026-07-16/issue-ledger.xlsx
inventory_roots: [FinalDelivery_2026-07-16]
canonical_poster_roots:
  - path: FinalDelivery_2026-07-16
    recursive: false
expected:
  release_movies: 2
  quarantine_movies: 1
  tier_counts: {S: 1, A: 1}
  pilot_movies: 1
  audit_rows: 1
  poster_states:
    machine_passed: 1
    content_conflict: 0
    placeholder: 0
    missing: 1
""",
        encoding="utf-8",
    )
    (config_dir / "gcp.yaml").write_text(
        """project_id: test-project
region: us-central1
required_account: test@example.com
archive_bucket_template: test-{project_id}
archive_prefix: source/v0.1
required_services: [storage.googleapis.com]
""",
        encoding="utf-8",
    )
    settings = Settings(
        repo_root=root.resolve(),
        release_version="v0.1",
        source_manifest=source_manifest.resolve(),
        source_validation=placeholders["validation.json"].resolve(),
        movies_csv=placeholders["movies.csv"].resolve(),
        movies_xlsx=placeholders["movies.xlsx"].resolve(),
        tiers_xlsx=placeholders["tiers.xlsx"].resolve(),
        issue_ledger_xlsx=issue_ledger.resolve(),
        inventory_roots=(source_dir.resolve(),),
        canonical_poster_roots=(),
        expected=ExpectedCounts(
            release_movies=2,
            quarantine_movies=1,
            tier_counts={"S": 1, "A": 1},
            pilot_movies=1,
            audit_rows=1,
            poster_states={
                "machine_passed": 1,
                "content_conflict": 0,
                "placeholder": 0,
                "missing": 1,
            },
        ),
        gcp=GcpSettings(
            project_id="test-project",
            region="us-central1",
            required_account="test@example.com",
            archive_bucket_template="test-{project_id}",
            archive_prefix="source/v0.1",
            required_services=("storage.googleapis.com",),
        ),
    )
    return ReleaseFixture(
        root=root,
        settings=settings,
        release_dir=release_dir,
        manifest_path=release_dir / "release_manifest.json",
        source_manifest=source_manifest,
        issue_ledger=issue_ledger,
        artifact_bytes={
            filename: (release_dir / filename).read_bytes()
            for filename in (
                "enrichment_provenance.parquet",
                "movie_tiers.parquet",
                "movies.csv",
                "movies.parquet",
                "pilot_movies.parquet",
                "poster_manifest.parquet",
            )
        },
    )


def _restore_release_artifacts(fixture: ReleaseFixture) -> None:
    for filename, content in fixture.artifact_bytes.items():
        (fixture.release_dir / filename).write_bytes(content)


def _movie_rows() -> list[dict[str, str]]:
    return [
        {
            "movie_id": "1970_GSQ_001",
            "chinese_title": "高山青",
            "english_title": "The Evergreen Mountains",
            "release_date": "1970-01-01",
            "production_region": "Hong Kong",
            "director": "Director A",
            "screenwriter": "Writer A",
            "cast": "Cast A",
            "genre": "Drama",
            "runtime_minutes": "90",
            "production_company": "Company A",
            "data_source": "HKFA",
        },
        {
            "movie_id": "1970_YCCS_001",
            "chinese_title": "玉女嬉春",
            "english_title": "Spring",
            "release_date": "1970-01-02",
            "production_region": "Hong Kong",
            "director": "Director B",
            "screenwriter": "Writer B",
            "cast": "Cast B",
            "genre": "Comedy",
            "runtime_minutes": "88",
            "production_company": "Company B",
            "data_source": "HKFA",
        },
    ]


def _poster_rows() -> list[dict[str, object]]:
    base: dict[str, object] = {
        "asset_type": "poster",
        "source_object_uri": "1970_GSQ_001.jpg",
        "original_object_uri": "assets/posters/originals/1970_GSQ_001.jpg",
        "derived_object_uri": "assets/posters/derived/1970_GSQ_001.webp",
        "quarantine_object_uri": "",
        "content_sha256": "a" * 64,
        "detected_format": "JPEG",
        "detected_mime_type": "image/jpeg",
        "detected_extension": ".jpg",
        "width": 600,
        "height": 900,
        "color_mode": "RGB",
        "byte_length": 100,
        "identity_confidence": None,
        "source_url": None,
        "rights_status": "unknown",
        "version": 1,
        "is_primary": False,
        "publishable": False,
        "created_at": None,
        "reviewed_at": None,
    }
    passed = {
        **base,
        "asset_id": "poster:1970_GSQ_001:v1",
        "movie_id": "1970_GSQ_001",
        "quality_status": "machine_passed",
    }
    missing = {
        **base,
        "asset_id": "poster:1970_YCCS_001:v1",
        "movie_id": "1970_YCCS_001",
        "source_object_uri": "",
        "original_object_uri": "",
        "derived_object_uri": "",
        "content_sha256": "",
        "detected_format": "",
        "detected_mime_type": "",
        "detected_extension": "",
        "width": None,
        "height": None,
        "color_mode": "",
        "byte_length": None,
        "quality_status": "missing",
    }
    return [passed, missing]


def _poster_schema() -> pa.Schema:
    return pa.schema(
        [
            ("asset_id", pa.string()),
            ("movie_id", pa.string()),
            ("asset_type", pa.string()),
            ("source_object_uri", pa.string()),
            ("original_object_uri", pa.string()),
            ("derived_object_uri", pa.string()),
            ("quarantine_object_uri", pa.string()),
            ("content_sha256", pa.string()),
            ("detected_format", pa.string()),
            ("detected_mime_type", pa.string()),
            ("detected_extension", pa.string()),
            ("width", pa.int32()),
            ("height", pa.int32()),
            ("color_mode", pa.string()),
            ("byte_length", pa.int64()),
            ("quality_status", pa.string()),
            ("identity_confidence", pa.float64()),
            ("source_url", pa.string()),
            ("rights_status", pa.string()),
            ("version", pa.int32()),
            ("is_primary", pa.bool_()),
            ("publishable", pa.bool_()),
            ("created_at", pa.string()),
            ("reviewed_at", pa.string()),
        ]
    )


def _strings_schema(*names: str) -> pa.Schema:
    return pa.schema([(name, pa.string()) for name in names])


def _write_csv(path: Path, rows: list[dict[str, str]]) -> None:
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def _write_parquet(
    path: Path, rows: list[dict[str, object]], schema: pa.Schema
) -> None:
    pq.write_table(pa.Table.from_pylist(rows, schema=schema), path)


def _write_issue_ledger(path: Path, movie_ids: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "隔離記錄"
    sheet.append(["影片唯一ID"])
    for movie_id in movie_ids:
        sheet.append([movie_id])
    workbook.save(path)


def _write_source_movie_csv(path: Path) -> None:
    columns = [
        "影片唯一ID",
        "中文片名",
        "英文片名",
        "上映日期",
        "出品地區",
        "導演",
        "編劇",
        "主演",
        "影片類型",
        "片長",
        "出品公司",
        "數據來源",
    ]
    rows = [
        dict(zip(columns, row.values(), strict=True))
        for row in _movie_rows()
    ]
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def _write_source_movies_workbook(path: Path) -> None:
    columns = [
        "影片唯一ID",
        "中文片名",
        "上映日期",
        "field_name",
        "original_value",
        "filled_value",
        "source_label",
        "source_url",
        "source_record_id",
        "match_method",
        "confidence",
        "matched_imdb_id",
        "retrieved_at",
        "notes",
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "補全審計"
    sheet.append(columns)
    row = pq.read_table(
        path.parents[1] / "data" / "release" / "v0.1" / "enrichment_provenance.parquet"
    ).to_pylist()[0]
    sheet.append(
        [
            row["movie_id"],
            row["chinese_title"],
            row["release_date"],
            *[row[column] for column in columns[3:]],
        ]
    )
    workbook.save(path)


def _write_source_tiers_workbook(path: Path) -> None:
    movie_columns = [
        "影片唯一ID",
        "中文片名",
        "英文片名",
        "上映日期",
        "出品地區",
        "導演",
        "編劇",
        "主演",
        "影片類型",
        "片長",
        "出品公司",
        "數據來源",
    ]
    tier_columns = [*movie_columns, "分級", "分級依據", "人工覆核", "試點影片"]
    workbook = Workbook()
    tier_sheet = workbook.active
    tier_sheet.title = "分級片單"
    tier_sheet.append(tier_columns)
    movies = _movie_rows()
    tiers = [("S", "reason", "required", True), ("A", "reason", "not_required", False)]
    for movie, tier in zip(movies, tiers, strict=True):
        tier_sheet.append([*movie.values(), *tier])
    pilot_columns = [
        *tier_columns,
        "手冊類型",
        "類型證據說明",
        "類型證據網址",
        "來源記錄ID",
        "權威層級",
    ]
    pilot_sheet = workbook.create_sheet("首批試點")
    pilot_sheet.append(pilot_columns)
    pilot_sheet.append(
        [
            *movies[0].values(),
            *tiers[0],
            "drama",
            "note",
            "https://example.com/evidence",
            "record-1",
            "primary",
        ]
    )
    workbook.save(path)


def _write_source_manifest(
    path: Path, issue_ledger: Path, placeholders: dict[str, Path]
) -> None:
    artifacts = [issue_ledger, *placeholders.values()]
    payload = {
        "version": "v0.1",
        "generated_at": "2026-07-17T04:08:41.762998+00:00",
        "validation_ok": True,
        "validation_checks": "22/22",
        "artifact_count": len(artifacts),
        "artifacts": [
            {
                "filename": artifact.name,
                "size_bytes": artifact.stat().st_size,
                "sha256": _sha256(artifact),
            }
            for artifact in sorted(artifacts, key=lambda item: item.name)
        ],
        "delivery_contract": {
            "release_rows": 2,
            "quarantine_rows": 1,
            "candidate_rows": 3,
            "tier_counts": {"S": 1, "A": 1},
            "pilot_count": 1,
            "audited_change_rows": 1,
        },
    }
    _write_manifest(path, payload)


def _rebind_source_manifest(fixture: ReleaseFixture) -> None:
    payload = json.loads(fixture.source_manifest.read_text(encoding="utf-8"))
    for artifact in payload["artifacts"]:
        if artifact["filename"] == fixture.issue_ledger.name:
            artifact["size_bytes"] = fixture.issue_ledger.stat().st_size
            artifact["sha256"] = _sha256(fixture.issue_ledger)
    _write_manifest(fixture.source_manifest, payload)


def _rebind_source_file(fixture: ReleaseFixture, path: Path) -> None:
    payload = json.loads(fixture.source_manifest.read_text(encoding="utf-8"))
    for artifact in payload["artifacts"]:
        if artifact["filename"] == path.name:
            artifact["size_bytes"] = path.stat().st_size
            artifact["sha256"] = _sha256(path)
            break
    else:
        raise AssertionError(f"source artifact is undeclared: {path.name}")
    _write_manifest(fixture.source_manifest, payload)


def _validation_report(version: str) -> dict[str, object]:
    return {
        "version": version,
        "ok": True,
        "checks_passed": 22,
        "checks_total": 22,
        "checks": [
            {
                "code": code,
                "ok": True,
                "severity": "blocking",
                "evidence": f"fixture:{code}",
            }
            for code in _VALIDATION_CODES
        ],
    }


def _mutate_bound_artifact(path: Path, mutation: str) -> None:
    content = path.read_bytes()
    if mutation == "delete":
        path.unlink()
    elif mutation == "replace":
        path.unlink()
        path.write_bytes(content)
    elif mutation == "in_place":
        changed = bytes([content[0] ^ 1]) + content[1:]
        path.write_bytes(changed)
    elif mutation == "hardlink":
        alias = path.parents[3] / f"hardlink-{path.name}"
        os.link(path, alias)
        changed = bytes([content[0] ^ 1]) + content[1:]
        alias.write_bytes(changed)
    else:
        raise AssertionError(f"unsupported mutation: {mutation}")


def _write_manifest(path: Path, payload: dict[str, object]) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        + "\n",
        encoding="utf-8",
    )


def _load_manifest(path: Path) -> dict[str, object]:
    return json.loads(path.read_text(encoding="utf-8"))


def _deepcopy(payload: dict[str, object]) -> dict[str, object]:
    return json.loads(json.dumps(payload))


def _artifact(payload: dict[str, object], filename: str) -> dict[str, object]:
    artifacts = payload["artifacts"]
    assert isinstance(artifacts, list)
    return next(
        item
        for item in artifacts
        if isinstance(item, dict) and str(item["relative_path"]).endswith(f"/{filename}")
    )


def _rebind_artifact(payload: dict[str, object], root: Path, filename: str) -> None:
    entry = _artifact(payload, filename)
    path = root / str(entry["relative_path"])
    entry["size_bytes"] = path.stat().st_size
    entry["sha256"] = _sha256(path)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
