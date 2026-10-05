"""Deterministic, fail-closed authority for canonical Release artifacts."""

from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import stat
from collections import Counter
from collections.abc import Iterator, Mapping, Sequence
from contextlib import contextmanager
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from jsonschema import Draft202012Validator, FormatChecker, ValidationError
from openpyxl import load_workbook

from .config import ConfigError, Settings
from .posters import PosterError, _atomic_replace_bytes, _locked_output_ancestors
from .release_data import (
    _PILOT_COLUMNS as SOURCE_PILOT_COLUMNS,
)
from .release_data import (
    _PROVENANCE_COLUMNS as SOURCE_PROVENANCE_COLUMNS,
)
from .release_data import (
    _TIER_COLUMNS as SOURCE_TIER_COLUMNS,
)
from .release_data import (
    MOVIE_COLUMNS,
    MOVIE_STORAGE_COLUMNS,
    _build_pilot_rows,
    _build_provenance_rows,
    _date_text,
    _text,
    join_tiers,
    validate_movie_rows,
)
from .release_lock import ReleaseLockError
from .release_lock import release_guard as _release_guard


class ManifestError(ValueError):
    """Raised when a Release cannot be authorized without ambiguity."""


@dataclass(frozen=True)
class ArtifactEntry:
    relative_path: str
    schema_path: str
    size_bytes: int
    sha256: str
    row_count: int

    def to_dict(self) -> dict[str, object]:
        return {
            "relative_path": self.relative_path,
            "schema_path": self.schema_path,
            "size_bytes": self.size_bytes,
            "sha256": self.sha256,
            "row_count": self.row_count,
        }


@dataclass(frozen=True)
class ReleaseManifest:
    release_version: str
    source_manifest_sha256: str
    source_generated_at: str
    expected_counts: Mapping[str, object]
    poster_state_counts: Mapping[str, int]
    artifacts: tuple[ArtifactEntry, ...]

    @property
    def artifact_count(self) -> int:
        return len(self.artifacts)

    def to_dict(self) -> dict[str, object]:
        return {
            "release_version": self.release_version,
            "source_manifest_sha256": self.source_manifest_sha256,
            "source_generated_at": self.source_generated_at,
            "expected_counts": dict(self.expected_counts),
            "poster_state_counts": dict(self.poster_state_counts),
            "artifacts": [artifact.to_dict() for artifact in self.artifacts],
        }


@dataclass(frozen=True)
class VerificationResult:
    valid: bool
    release_version: str
    artifact_count: int
    movie_count: int


@dataclass(frozen=True)
class _FileBinding:
    device: int
    inode: int
    size_bytes: int
    modified_ns: int
    changed_ns: int
    sha256: str


@dataclass
class _ManifestPublicationState:
    previous_manifest: bytes | None = None
    candidate_binding: _FileBinding | None = None
    candidate_content: bytes | None = None
    installed: bool = False


@dataclass(frozen=True)
class _HeldFile:
    root: Path
    path: Path
    label: str
    descriptor: int
    content: bytes
    binding: _FileBinding


_MOVIE_SCHEMA = "schemas/movie.schema.json"
_MEDIA_SCHEMA = "schemas/media_asset.schema.json"
_MANIFEST_SCHEMA = "schemas/release_manifest.schema.json"
_ARTIFACT_SCHEMAS = {
    "enrichment_provenance.parquet": _MOVIE_SCHEMA,
    "movie_tiers.parquet": _MOVIE_SCHEMA,
    "movies.csv": _MOVIE_SCHEMA,
    "movies.parquet": _MOVIE_SCHEMA,
    "pilot_movies.parquet": _MOVIE_SCHEMA,
    "poster_manifest.parquet": _MEDIA_SCHEMA,
}
_TIER_COLUMNS = ["movie_id", "tier", "tier_reason", "human_review", "pilot_movie"]
_PILOT_COLUMNS = [
    "movie_id",
    "handbook_genre",
    "evidence_note",
    "evidence_url",
    "source_record_id",
]
_PROVENANCE_COLUMNS = [
    "movie_id",
    "chinese_title",
    "release_date",
    "field_name",
    "original_value",
    "filled_value",
    "source_label",
    "source_url",
    "source_record_id",
    "match_method",
    "confidence",
    "matched_imdb_id",
    "retrieved_at",
    "notes",
]
_POSTER_COLUMNS = [
    "asset_id",
    "movie_id",
    "asset_type",
    "source_object_uri",
    "original_object_uri",
    "derived_object_uri",
    "quarantine_object_uri",
    "content_sha256",
    "detected_format",
    "detected_mime_type",
    "detected_extension",
    "width",
    "height",
    "color_mode",
    "byte_length",
    "quality_status",
    "identity_confidence",
    "source_url",
    "rights_status",
    "version",
    "is_primary",
    "publishable",
    "created_at",
    "reviewed_at",
]
_UNIQUE_CHILDREN = {
    "movie_tiers.parquet": ("movie_id",),
    "pilot_movies.parquet": ("movie_id",),
    "poster_manifest.parquet": ("asset_id", "movie_id"),
    "enrichment_provenance.parquet": ("movie_id", "field_name", "source_record_id"),
}
_V12_PLACEHOLDER_SHA256 = "2f17500b570bc2cc06e1f743fc1d84bcfb4ca89c2a10fc666ab81cc987c064b8"
_V12_POSTER_SEMANTIC_SHA256 = "bd5418911cc743a745013875287720ffc2fabc5b27935664a48b45343924c15a"
_V12_POSTER_COUNTS = {
    "machine_passed": 4545,
    "content_conflict": 73,
    "placeholder": 9,
    "missing": 31,
}
_V12_SOURCE_MANIFEST_SHA256 = (
    "cbd529e126e8b72dccac302d9746fcfa63318caecbcda24548c05e282e13f974"
)
_V12_SOURCE_GENERATED_AT = "2026-07-17T04:08:41.762998+00:00"
_V12_EXPECTED_COUNTS = {
    "release_movies": 4658,
    "quarantine_movies": 1264,
    "tier_counts": {"S": 50, "A": 313, "B": 4295},
    "pilot_movies": 24,
    "audit_rows": 1008,
}
_V12_CANONICAL_SEMANTIC_SHA256 = {
    "enrichment_provenance.parquet": "04e5b5712c7f8ee04e7696ee08b5d2821a6108759948c134b3ef237d57506ba0",
    "movie_tiers.parquet": "0070331c8ed428f4702f698ba94a653c8388183803ee632f37a5214deeb755c9",
    "movies.csv": "a0ce946130ff9e53042004ff9ca68a90b399fc81c949d5896df799a5c40dd567",
    "movies.parquet": "a0ce946130ff9e53042004ff9ca68a90b399fc81c949d5896df799a5c40dd567",
    "pilot_movies.parquet": "4d35ee6df54a018000bafa7efde29d6fa53c3efba4bde20a5f3a6e22b2c3f862",
}
_V12_VALIDATION_CODES = (
    "COLUMN_CONTRACT",
    "RELEASE_ROW_COUNT",
    "RELEASE_ID_UNIQUE",
    "REQUIRED_FIELDS_COMPLETE",
    "DATE_FORMAT_AND_RANGE",
    "RUNTIME_POSITIVE_INTEGER",
    "NO_SENTINELS",
    "CSV_XLSX_EXACT_ALIGNMENT",
    "DATA_TIER_METADATA_ALIGNMENT",
    "SUPPLEMENT_AUDIT_COMPLETE",
    "AUDIT_SOURCE_TRACEABLE",
    "AUDIT_ISSUE_LEDGER_ALIGNMENT",
    "TIER_COUNTS",
    "PILOT_ALIGNMENT",
    "S_TIER_COMPLETE",
    "RELEASE_QUARANTINE_DISJOINT",
    "CANDIDATE_RECONCILIATION",
    "SOURCE_FIDELITY_TITLE_FORMAT",
    "WORKBOOK_ERROR_SCAN",
    "MERGED_LEDGER_ROW_COUNTS",
    "MERGED_LEDGER_ENTITY_ALIGNMENT",
    "MERGED_LEDGER_ENCODING_AND_DATE_PRESERVATION",
)


def build_release_manifest(settings: Settings) -> ReleaseManifest:
    """Validate all canonical inputs, then atomically publish the sole authority."""
    root = _require_repo_root(settings.repo_root)
    release_dir = _canonical_release_dir(root, settings.release_version)
    manifest_path = release_dir / "release_manifest.json"
    publication = _ManifestPublicationState()
    with _release_manifest_publication_guard(
        root, settings.release_version, manifest_path, publication
    ):
        _require_release_namespace(release_dir, allow_manifest=True)
        source = _load_source_evidence(settings, root)
        with _hold_release_artifacts(root, release_dir) as held_artifacts:
            artifact_bytes = {
                filename: held.content for filename, held in held_artifacts.items()
            }
            artifact_bindings = {
                filename: held.binding for filename, held in held_artifacts.items()
            }
            rows = _parse_release_rows(artifact_bytes, root)
            counts = _verify_release_content(rows, settings, source)

            entries = tuple(
                ArtifactEntry(
                    relative_path=(
                        Path("data")
                        / "release"
                        / settings.release_version
                        / filename
                    ).as_posix(),
                    schema_path=_ARTIFACT_SCHEMAS[filename],
                    size_bytes=len(artifact_bytes[filename]),
                    sha256=_sha256_bytes(artifact_bytes[filename]),
                    row_count=len(rows[filename]),
                )
                for filename in sorted(_ARTIFACT_SCHEMAS)
            )
            manifest = ReleaseManifest(
                release_version=settings.release_version,
                source_manifest_sha256=source.manifest_sha256,
                source_generated_at=source.generated_at,
                expected_counts=counts,
                poster_state_counts={
                    state: Counter(
                        str(row["quality_status"])
                        for row in rows["poster_manifest.parquet"]
                    )[state]
                    for state in sorted(settings.expected.poster_states)
                },
                artifacts=entries,
            )
            payload = manifest.to_dict()
            _validate_manifest_schema(payload, root, manifest_path)
            content = (
                json.dumps(
                    payload,
                    ensure_ascii=False,
                    sort_keys=True,
                    separators=(",", ":"),
                )
                + "\n"
            ).encode("utf-8")
            previous_manifest = _read_previous_manifest(root, manifest_path)
            publication.previous_manifest = previous_manifest
            candidate_binding: _FileBinding | None = None
            candidate_content: bytes | None = None

            def verify_installed_candidate() -> None:
                nonlocal candidate_binding, candidate_content
                with _hold_file(
                    root,
                    manifest_path,
                    "release manifest",
                    deny_write_delete=False,
                ) as held_candidate:
                    candidate_content = held_candidate.content
                    candidate_binding = held_candidate.binding
                    publication.candidate_content = candidate_content
                    publication.candidate_binding = candidate_binding
                    if held_candidate.content != content:
                        raise ManifestError(
                            "installed release manifest content changed"
                        )
                    _verify_held_file(held_candidate)
                _verify_installed_release_authority(
                    root,
                    release_dir,
                    artifact_bindings,
                    artifact_bytes,
                )
                _verify_held_files(held_artifacts)

            try:
                _atomic_replace_bytes(
                    root,
                    manifest_path,
                    content,
                    pre_install=lambda: _verify_held_files(held_artifacts),
                    post_install=verify_installed_candidate,
                )
                publication.installed = True
            except (ManifestError, PosterError, OSError) as exc:
                try:
                    _rollback_release_manifest(
                        root,
                        manifest_path,
                        previous_manifest,
                        failed_binding=candidate_binding,
                        failed_content=candidate_content,
                    )
                except (ManifestError, PosterError) as restore_exc:
                    raise ManifestError(
                        "cannot restore prior release manifest after failed publication"
                    ) from restore_exc
                if isinstance(exc, ManifestError):
                    raise
                if isinstance(exc, OSError):
                    raise ManifestError("artifact changed after binding") from exc
                raise ManifestError(f"cannot publish release manifest: {exc}") from exc
            return manifest


@contextmanager
def _release_manifest_publication_guard(
    root: Path,
    release_version: str,
    manifest_path: Path,
    publication: _ManifestPublicationState,
) -> Iterator[None]:
    """Rollback an installed candidate before a guard-exit failure escapes."""
    try:
        with _release_guard(root, release_version, exclusive=True):
            yield
    except ReleaseLockError:
        if publication.installed:
            try:
                with _release_guard(root, release_version, exclusive=True):
                    _rollback_release_manifest(
                        root,
                        manifest_path,
                        publication.previous_manifest,
                        failed_binding=publication.candidate_binding,
                        failed_content=publication.candidate_content,
                    )
            except ReleaseLockError:
                raise
            except (ManifestError, PosterError, OSError) as restore_exc:
                raise ManifestError(
                    "cannot restore prior release manifest after Release lock failure"
                ) from restore_exc
            publication.installed = False
        raise


def verify_release_manifest(path: Path) -> VerificationResult:
    """Verify a canonical manifest in the documented fail-closed order."""
    manifest_path, root = _locate_manifest(path)
    locked_version = manifest_path.parent.name
    with _release_guard(root, locked_version, exclusive=False):
        return _verify_release_manifest_locked(manifest_path, root, locked_version)


def _verify_release_manifest_locked(
    manifest_path: Path, root: Path, locked_version: str
) -> VerificationResult:
    """Verify while the manifest and all six artifact handles remain held."""

    # 1. Schema validity: strict JSON, canonical shape/path set, and valid schemas.
    with _hold_file(root, manifest_path, "release manifest") as held_manifest:
        manifest_content = held_manifest.content
        payload = _read_json_bytes(manifest_content, "release manifest")
        _validate_manifest_schema(payload, root, manifest_path)
        release_version = _required_string(payload, "release_version", "release manifest")
        if release_version != locked_version:
            raise ManifestError("source-manifest identity mismatch: release version")

        # 2. Source identity, including every configured source evidence byte.
        try:
            settings = Settings.load(root)
        except (ConfigError, OSError, ValueError) as exc:
            raise ManifestError(f"source-manifest identity failed: {exc}") from exc
        source = _load_source_evidence(settings, root)
        if release_version != settings.release_version:
            raise ManifestError("source-manifest identity mismatch: release version")
        if payload["source_manifest_sha256"] != source.manifest_sha256:
            raise ManifestError("source-manifest identity mismatch: SHA-256")
        if payload["source_generated_at"] != source.generated_at:
            raise ManifestError("source-manifest identity mismatch: generated_at")
        _require_declared_counts(payload, source)

        artifacts = _manifest_artifacts(payload)

        # 3. Artifact existence and canonical namespace completeness.
        release_dir = _canonical_release_dir(root, release_version)
        _require_release_namespace(
            release_dir, allow_manifest=True, require_manifest=True
        )

        # 4. Bind all six simultaneously; handles stay open through valid=True.
        with _hold_release_artifacts(root, release_dir) as held_artifacts:
            artifact_bytes: dict[str, bytes] = {}
            artifact_bindings: dict[str, _FileBinding] = {}
            for filename in sorted(artifacts):
                held = held_artifacts[filename]
                content = held.content
                entry = artifacts[filename]
                if _sha256_bytes(content) != entry["sha256"]:
                    raise ManifestError(f"artifact SHA-256 mismatch: {filename}")
                if len(content) != entry["size_bytes"]:
                    raise ManifestError(f"artifact size mismatch: {filename}")
                artifact_bytes[filename] = content
                artifact_bindings[filename] = held.binding

            rows = _parse_release_rows(artifact_bytes, root)

    # 5. Row-count reconciliation.
            for filename in sorted(artifacts):
                if len(rows[filename]) != artifacts[filename]["row_count"]:
                    raise ManifestError(f"artifact row-count mismatch: {filename}")
            _verify_row_counts(rows, payload)

    # 6. Unique identifiers and agreement of the two movie serializations.
            release_ids = _verify_unique_ids(rows)

    # 7. Every child artifact is keyed only to the formal Release parent table.
            _verify_foreign_keys(rows, release_ids)
            _verify_pilot_tier_reconciliation(rows)
            _verify_source_reconciliation(rows, source.canonical_rows)

    # 8. Exactly one expected poster state exists for every movie.
            _verify_poster_reconciliation(rows, release_ids, payload, settings)

    # 9. Quarantine identifiers are disjoint from formal Release identifiers.
            try:
                quarantine_ids, quarantine_count = _read_quarantine_ids(
                    source.source_bytes[settings.issue_ledger_xlsx.name]
                )
            except OSError as exc:
                raise ManifestError(
                    "artifact or release manifest changed after binding"
                ) from exc
            if quarantine_count != settings.expected.quarantine_movies:
                raise ManifestError(
                    "Quarantine disjointness failed: "
                    f"rows={quarantine_count}, expected={settings.expected.quarantine_movies}"
                )
            overlap = release_ids & quarantine_ids
            if overlap:
                raise ManifestError(
                    f"Quarantine disjointness failed: overlap={sorted(overlap)[:5]}"
                )
            _rebind_release_artifacts(
                root,
                release_dir,
                artifact_bindings,
                artifact_bytes,
                require_manifest=True,
                manifest_binding=held_manifest.binding,
                manifest_content=manifest_content,
            )
            _verify_held_files(held_artifacts)
            _verify_held_file(held_manifest)
            return VerificationResult(
                valid=True,
                release_version=release_version,
                artifact_count=len(artifacts),
                movie_count=len(release_ids),
            )


@dataclass(frozen=True)
class _SourceEvidence:
    manifest_sha256: str
    generated_at: str
    expected_counts: Mapping[str, object]
    source_bytes: Mapping[str, bytes]
    canonical_rows: Mapping[str, list[dict[str, object]]]


def _load_source_evidence(settings: Settings, root: Path) -> _SourceEvidence:
    try:
        content = _read_safe_bytes(root, settings.source_manifest, "source manifest")
        manifest = _read_json_bytes(content, "source manifest")
        version = _required_string(manifest, "version", "source manifest")
        generated_at = _required_string(manifest, "generated_at", "source manifest")
        if version != settings.release_version:
            raise ManifestError("source-manifest identity mismatch: release version")
        try:
            parsed_time = datetime.fromisoformat(generated_at)
        except ValueError as exc:
            raise ManifestError("source-manifest identity has invalid generated_at") from exc
        if parsed_time.tzinfo is None:
            raise ManifestError("source-manifest identity generated_at must include timezone")
        artifacts = manifest.get("artifacts")
        if not isinstance(artifacts, list):
            raise ManifestError("source-manifest identity artifacts must be a list")
        declared: dict[str, Mapping[str, object]] = {}
        for item in artifacts:
            if not isinstance(item, dict):
                raise ManifestError("source-manifest identity artifact is invalid")
            filename = item.get("filename")
            if not isinstance(filename, str) or Path(filename).name != filename:
                raise ManifestError("source-manifest identity filename is invalid")
            if filename in declared:
                raise ManifestError("source-manifest identity has duplicate artifacts")
            declared[filename] = item
        artifact_count = manifest.get("artifact_count")
        if (
            not isinstance(artifact_count, int)
            or isinstance(artifact_count, bool)
            or artifact_count != len(artifacts)
        ):
            raise ManifestError("source-manifest identity artifact_count mismatch")
        if (
            manifest.get("validation_ok") is not True
            or manifest.get("validation_checks") != "22/22"
        ):
            raise ManifestError("source validation must be ok=true and 22/22")
        source_bytes: dict[str, bytes] = {}
        source_directory = settings.source_manifest.parent
        for filename, item in sorted(declared.items()):
            expected_size = item.get("size_bytes")
            expected_hash = item.get("sha256")
            if (
                not isinstance(expected_size, int)
                or isinstance(expected_size, bool)
                or expected_size < 0
                or not isinstance(expected_hash, str)
                or len(expected_hash) != 64
            ):
                raise ManifestError(
                    f"source-manifest identity metadata invalid: {filename}"
                )
            source_path = source_directory / filename
            source_content = _read_safe_bytes(root, source_path, f"source {filename}")
            if (
                len(source_content) != expected_size
                or _sha256_bytes(source_content) != expected_hash
            ):
                raise ManifestError(f"source-manifest identity mismatch: {filename}")
            source_bytes[filename] = source_content

        configured = (
            settings.source_validation,
            settings.movies_csv,
            settings.movies_xlsx,
            settings.tiers_xlsx,
            settings.issue_ledger_xlsx,
        )
        for source_path in configured:
            item = declared.get(source_path.name)
            if item is None:
                raise ManifestError(
                    f"source-manifest identity missing artifact: {source_path.name}"
                )
            if source_path != source_directory / source_path.name:
                raise ManifestError(
                    f"source-manifest identity path mismatch: {source_path.name}"
                )
        expected_counts = _source_expected_counts(manifest, settings)
        _validate_source_validation(
            source_bytes[settings.source_validation.name], settings.release_version
        )
        canonical_rows = _derive_canonical_source_rows(settings, source_bytes)
        manifest_sha256 = _sha256_bytes(content)
        _verify_v12_frozen_authority(
            release_version=settings.release_version,
            source_manifest_sha256=manifest_sha256,
            source_generated_at=generated_at,
            expected_counts=expected_counts,
            canonical_digests={
                filename: _poster_semantic_sha256(rows)
                for filename, rows in canonical_rows.items()
            },
        )
        return _SourceEvidence(
            manifest_sha256=manifest_sha256,
            generated_at=generated_at,
            expected_counts=expected_counts,
            source_bytes=source_bytes,
            canonical_rows=canonical_rows,
        )
    except ManifestError as exc:
        if str(exc).startswith("source-manifest identity"):
            raise
        raise ManifestError(f"source-manifest identity failed: {exc}") from exc


def _source_expected_counts(
    manifest: Mapping[str, object], settings: Settings
) -> Mapping[str, object]:
    contract = manifest.get("delivery_contract")
    if not isinstance(contract, dict):
        raise ManifestError("source-manifest identity delivery contract is missing")
    fields = {
        "release_movies": contract.get("release_rows"),
        "quarantine_movies": contract.get("quarantine_rows"),
        "tier_counts": contract.get("tier_counts"),
        "pilot_movies": contract.get("pilot_count"),
        "audit_rows": contract.get("audited_change_rows"),
    }
    expected = {
        "release_movies": settings.expected.release_movies,
        "quarantine_movies": settings.expected.quarantine_movies,
        "tier_counts": dict(settings.expected.tier_counts),
        "pilot_movies": settings.expected.pilot_movies,
        "audit_rows": settings.expected.audit_rows,
    }
    if fields != expected:
        raise ManifestError("source-manifest identity delivery counts mismatch")
    candidate = contract.get("candidate_rows")
    if candidate != settings.expected.release_movies + settings.expected.quarantine_movies:
        raise ManifestError("source-manifest identity candidate count mismatch")
    return expected


def _validate_source_validation(content: bytes, release_version: str) -> None:
    report = _read_json_bytes(content, "source validation")
    checks = report.get("checks")
    if (
        report.get("ok") is not True
        or report.get("checks_passed") != 22
        or report.get("checks_total") != 22
        or not isinstance(checks, list)
        or len(checks) != 22
        or any(not isinstance(check, dict) or check.get("ok") is not True for check in checks)
    ):
        raise ManifestError("source validation must be ok=true and 22/22")
    assert isinstance(checks, list)
    if (
        report.get("version") != release_version
        or tuple(check.get("code") for check in checks) != _V12_VALIDATION_CODES
        or any(
            set(check) != {"code", "ok", "severity", "evidence"}
            or check.get("severity") != "blocking"
            or not isinstance(check.get("evidence"), str)
            or not check.get("evidence")
            for check in checks
        )
    ):
        raise ManifestError("source validation failed: expected exact 22 identities")


def _verify_v12_frozen_authority(
    *,
    release_version: str,
    source_manifest_sha256: str,
    source_generated_at: str,
    expected_counts: Mapping[str, object],
    canonical_digests: Mapping[str, str],
) -> None:
    if release_version != "v1.2":
        return
    if (
        source_manifest_sha256 != _V12_SOURCE_MANIFEST_SHA256
        or source_generated_at != _V12_SOURCE_GENERATED_AT
        or dict(expected_counts) != _V12_EXPECTED_COUNTS
    ):
        raise ManifestError("source-manifest identity mismatch: frozen source manifest")
    if dict(canonical_digests) != _V12_CANONICAL_SEMANTIC_SHA256:
        raise ManifestError("source-manifest identity mismatch: frozen canonical semantics")


def _derive_canonical_source_rows(
    settings: Settings, source_bytes: Mapping[str, bytes]
) -> Mapping[str, list[dict[str, object]]]:
    try:
        source_movies = _read_source_movie_rows(
            source_bytes[settings.movies_csv.name]
        )
        validate_movie_rows(source_movies, columns=MOVIE_COLUMNS)
        source_movies.sort(key=lambda row: _text(row["影片唯一ID"]))
        storage_movies: list[dict[str, object]] = []
        column_map = dict(zip(MOVIE_COLUMNS, MOVIE_STORAGE_COLUMNS, strict=True))
        for row in source_movies:
            storage_row: dict[str, object] = {
                column_map[column]: _text(row[column]) for column in MOVIE_COLUMNS
            }
            storage_row["release_date"] = _date_text(row["上映日期"])
            storage_movies.append(storage_row)

        tier_source = _read_source_sheet_rows(
            source_bytes[settings.tiers_xlsx.name], "分級片單", SOURCE_TIER_COLUMNS
        )
        tier_records = join_tiers(source_movies, tier_source)
        storage_tiers: list[dict[str, object]] = [
            {
                "movie_id": record.movie_id,
                "tier": record.tier,
                "tier_reason": record.tier_reason,
                "human_review": record.human_review,
                "pilot_movie": record.pilot_movie,
            }
            for record in tier_records
        ]
        release_ids = {str(row["movie_id"]) for row in storage_movies}
        pilot_source = _read_source_sheet_rows(
            source_bytes[settings.tiers_xlsx.name], "首批試點", SOURCE_PILOT_COLUMNS
        )
        storage_pilots = _build_pilot_rows(
            pilot_source, tier_records, release_ids
        )
        provenance_source = _read_source_sheet_rows(
            source_bytes[settings.movies_xlsx.name],
            "補全審計",
            SOURCE_PROVENANCE_COLUMNS,
        )
        storage_provenance = _build_provenance_rows(
            provenance_source, release_ids
        )
        return {
            "movies.csv": storage_movies,
            "movies.parquet": storage_movies,
            "movie_tiers.parquet": storage_tiers,
            "pilot_movies.parquet": storage_pilots,
            "enrichment_provenance.parquet": storage_provenance,
        }
    except ManifestError:
        raise
    except Exception as exc:
        raise ManifestError(f"source reconciliation failed: {exc}") from exc


def _read_source_movie_rows(content: bytes) -> list[dict[str, object]]:
    try:
        reader = csv.DictReader(io.StringIO(content.decode("utf-8-sig"), newline=""))
        if reader.fieldnames != MOVIE_COLUMNS:
            raise ManifestError("source reconciliation failed: movie columns")
        rows: list[dict[str, object]] = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ManifestError("source reconciliation failed: movie row")
            rows.append(dict(row))
        return rows
    except ManifestError:
        raise
    except (UnicodeError, csv.Error) as exc:
        raise ManifestError("source reconciliation failed: movie CSV unreadable") from exc


def _read_source_sheet_rows(
    content: bytes, sheet_name: str, columns: Sequence[str]
) -> list[dict[str, object]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if sheet_name not in workbook.sheetnames:
            raise ManifestError(f"source reconciliation failed: missing sheet {sheet_name}")
        iterator = workbook[sheet_name].iter_rows(values_only=True)
        header = list(next(iterator, ()))
        if header != list(columns):
            raise ManifestError(
                f"source reconciliation failed: {sheet_name} columns"
            )
        return [
            {column: row[index] for index, column in enumerate(columns)}
            for row in iterator
        ]
    except ManifestError:
        raise
    except Exception as exc:
        raise ManifestError(
            f"source reconciliation failed: {sheet_name} unreadable"
        ) from exc


def _read_and_bind_artifacts(
    root: Path, release_version: str
) -> tuple[dict[str, bytes], dict[str, _FileBinding]]:
    release_dir = _canonical_release_dir(root, release_version)
    content: dict[str, bytes] = {}
    bindings: dict[str, _FileBinding] = {}
    for filename in sorted(_ARTIFACT_SCHEMAS):
        path = release_dir / filename
        try:
            content[filename], bindings[filename] = _read_bound_file(
                root, path, f"artifact {filename}"
            )
        except ManifestError as exc:
            if "does not exist" in str(exc):
                raise ManifestError(f"authoritative artifact missing: {filename}") from exc
            raise
    return content, bindings


@contextmanager
def _hold_release_artifacts(
    root: Path, release_dir: Path
) -> Iterator[dict[str, _HeldFile]]:
    """Open and retain one simultaneous six-artifact authorization snapshot."""
    held: dict[str, _HeldFile] = {}
    try:
        for filename in sorted(_ARTIFACT_SCHEMAS):
            path = release_dir / filename
            try:
                held[filename] = _open_held_file(
                    root, path, f"artifact {filename}", deny_write_delete=True
                )
            except ManifestError as exc:
                if "does not exist" in str(exc):
                    raise ManifestError(
                        f"authoritative artifact missing: {filename}"
                    ) from exc
                raise
        yield held
    finally:
        for item in reversed(tuple(held.values())):
            os.close(item.descriptor)


@contextmanager
def _hold_file(
    root: Path,
    path: Path,
    label: str,
    *,
    deny_write_delete: bool = True,
) -> Iterator[_HeldFile]:
    held = _open_held_file(
        root, path, label, deny_write_delete=deny_write_delete
    )
    try:
        yield held
    finally:
        os.close(held.descriptor)


def _open_held_file(
    root: Path, path: Path, label: str, *, deny_write_delete: bool
) -> _HeldFile:
    before = _require_safe_regular_file(root, path, label)
    descriptor: int | None = None
    try:
        descriptor = _open_protected_descriptor(
            path, deny_write_delete=deny_write_delete
        )
        opened = os.fstat(descriptor)
        if (
            not stat.S_ISREG(opened.st_mode)
            or _is_reparse(opened)
            or (opened.st_dev, opened.st_ino) != (before.st_dev, before.st_ino)
        ):
            raise ManifestError(f"{label} changed during protected open")
        content = _read_descriptor(descriptor)
        after = os.fstat(descriptor)
        rebound = _require_safe_regular_file(root, path, label)
        if (
            (after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns)
            != (opened.st_dev, opened.st_ino, opened.st_size, opened.st_mtime_ns)
            or (rebound.st_dev, rebound.st_ino) != (opened.st_dev, opened.st_ino)
        ):
            raise ManifestError(f"{label} changed during protected open")
        binding = _binding_from_stat(rebound, content)
        held = _HeldFile(root, path, label, descriptor, content, binding)
        descriptor = None
        return held
    except ManifestError:
        raise
    except OSError as exc:
        raise ManifestError(f"{label} is unreadable") from exc
    finally:
        if descriptor is not None:
            os.close(descriptor)


def _open_protected_descriptor(path: Path, *, deny_write_delete: bool) -> int:
    if os.name != "nt":
        import fcntl

        descriptor = os.open(
            path,
            os.O_RDONLY | getattr(os, "O_CLOEXEC", 0) | getattr(os, "O_NOFOLLOW", 0),
        )
        try:
            fcntl.flock(  # type: ignore[attr-defined]
                descriptor,
                fcntl.LOCK_SH,  # type: ignore[attr-defined]
            )
        except OSError:
            os.close(descriptor)
            raise
        return descriptor

    import ctypes
    import msvcrt

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [
        ctypes.c_wchar_p,
        ctypes.c_uint32,
        ctypes.c_uint32,
        ctypes.c_void_p,
        ctypes.c_uint32,
        ctypes.c_uint32,
        ctypes.c_void_p,
    ]
    create_file.restype = ctypes.c_void_p
    handle = create_file(
        str(path),
        0x80000000 | 0x80,  # GENERIC_READ | FILE_READ_ATTRIBUTES
        0x1 if deny_write_delete else 0x7,
        # FILE_SHARE_READ only for authority; share all for the install handle callback.
        None,
        3,  # OPEN_EXISTING
        0x00200000,  # FILE_FLAG_OPEN_REPARSE_POINT
        None,
    )
    invalid_handle = ctypes.c_void_p(-1).value
    if handle == invalid_handle:
        error = ctypes.get_last_error()
        raise OSError(error, os.strerror(error))
    try:
        return msvcrt.open_osfhandle(int(handle), os.O_BINARY | os.O_RDONLY)
    except OSError:
        kernel32.CloseHandle(ctypes.c_void_p(handle))
        raise


def _read_descriptor(descriptor: int) -> bytes:
    os.lseek(descriptor, 0, os.SEEK_SET)
    chunks: list[bytes] = []
    while True:
        chunk = os.read(descriptor, 1024 * 1024)
        if not chunk:
            return b"".join(chunks)
        chunks.append(chunk)


def _binding_from_stat(path_stat: os.stat_result, content: bytes) -> _FileBinding:
    return _FileBinding(
        device=path_stat.st_dev,
        inode=path_stat.st_ino,
        size_bytes=path_stat.st_size,
        modified_ns=path_stat.st_mtime_ns,
        changed_ns=path_stat.st_ctime_ns,
        sha256=_sha256_bytes(content),
    )


def _verify_held_files(held_files: Mapping[str, _HeldFile]) -> None:
    for filename in sorted(held_files):
        _verify_held_file(held_files[filename])


def _verify_held_file(held: _HeldFile) -> None:
    """Prove a held file and its canonical name still denote the bound bytes."""
    try:
        content = _read_descriptor(held.descriptor)
        current = os.fstat(held.descriptor)
        if (
            (current.st_dev, current.st_ino, current.st_size, current.st_mtime_ns)
            != (
                held.binding.device,
                held.binding.inode,
                held.binding.size_bytes,
                held.binding.modified_ns,
            )
            or content != held.content
        ):
            raise ManifestError(f"{held.label} changed after binding")
        named = _require_safe_regular_file(held.root, held.path, held.label)
        if _binding_from_stat(named, content) != held.binding:
            raise ManifestError(f"{held.label} changed after binding")
    except ManifestError:
        raise
    except OSError as exc:
        raise ManifestError(f"{held.label} changed after binding") from exc


def _parse_release_rows(
    artifact_bytes: Mapping[str, bytes], root: Path
) -> dict[str, list[dict[str, object]]]:
    movie_schema = _read_schema(root, _MOVIE_SCHEMA)
    media_schema = _read_schema(root, _MEDIA_SCHEMA)
    movie_validator = Draft202012Validator(movie_schema, format_checker=FormatChecker())
    media_validator = Draft202012Validator(media_schema, format_checker=FormatChecker())
    rows: dict[str, list[dict[str, object]]] = {}
    rows["movies.csv"] = _read_csv_rows(artifact_bytes["movies.csv"])
    rows["movies.parquet"] = _read_parquet_rows(
        artifact_bytes["movies.parquet"], "movies.parquet", MOVIE_STORAGE_COLUMNS
    )
    rows["movie_tiers.parquet"] = _read_parquet_rows(
        artifact_bytes["movie_tiers.parquet"], "movie_tiers.parquet", _TIER_COLUMNS
    )
    rows["pilot_movies.parquet"] = _read_parquet_rows(
        artifact_bytes["pilot_movies.parquet"], "pilot_movies.parquet", _PILOT_COLUMNS
    )
    rows["enrichment_provenance.parquet"] = _read_parquet_rows(
        artifact_bytes["enrichment_provenance.parquet"],
        "enrichment_provenance.parquet",
        _PROVENANCE_COLUMNS,
    )
    rows["poster_manifest.parquet"] = _read_parquet_rows(
        artifact_bytes["poster_manifest.parquet"], "poster_manifest.parquet", _POSTER_COLUMNS
    )
    for filename in ("movies.csv", "movies.parquet"):
        for index, row in enumerate(rows[filename], start=1):
            try:
                movie_validator.validate(row)
            except ValidationError as exc:
                raise ManifestError(
                    f"artifact data schema invalid: {filename} row {index}"
                ) from exc
    for index, row in enumerate(rows["poster_manifest.parquet"], start=1):
        try:
            media_validator.validate(row)
        except ValidationError as exc:
            raise ManifestError(
                f"artifact data schema invalid: poster_manifest.parquet row {index}"
            ) from exc
    return rows


def _read_csv_rows(content: bytes) -> list[dict[str, object]]:
    try:
        text = content.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text, newline=""))
        if reader.fieldnames != MOVIE_STORAGE_COLUMNS:
            raise ManifestError("artifact data schema invalid: movies.csv columns")
        rows: list[dict[str, object]] = []
        for row in reader:
            if None in row or any(value is None for value in row.values()):
                raise ManifestError("artifact data schema invalid: movies.csv row")
            rows.append(dict(row))
        return rows
    except (UnicodeError, csv.Error) as exc:
        raise ManifestError("artifact data schema invalid: movies.csv unreadable") from exc


def _read_parquet_rows(
    content: bytes, filename: str, expected_columns: Sequence[str]
) -> list[dict[str, object]]:
    try:
        table = pq.read_table(pa.BufferReader(content))
    except (OSError, pa.ArrowException) as exc:
        raise ManifestError(f"artifact data schema invalid: {filename} unreadable") from exc
    expected_schema = _expected_arrow_schema(filename)
    if (
        table.schema.names != list(expected_columns)
        or not table.schema.equals(expected_schema, check_metadata=False)
    ):
        raise ManifestError(f"artifact data schema invalid: {filename} Arrow schema")
    return table.to_pylist()


def _expected_arrow_schema(filename: str) -> pa.Schema:
    string_schema = lambda columns: pa.schema([(column, pa.string()) for column in columns])
    if filename == "movies.parquet":
        return string_schema(MOVIE_STORAGE_COLUMNS)
    if filename == "movie_tiers.parquet":
        return pa.schema(
            [
                ("movie_id", pa.string()),
                ("tier", pa.string()),
                ("tier_reason", pa.string()),
                ("human_review", pa.string()),
                ("pilot_movie", pa.bool_()),
            ]
        )
    if filename == "pilot_movies.parquet":
        return string_schema(_PILOT_COLUMNS)
    if filename == "enrichment_provenance.parquet":
        return string_schema(_PROVENANCE_COLUMNS)
    if filename == "poster_manifest.parquet":
        return pa.schema(
            [
                ("asset_id", pa.string()),
                ("movie_id", pa.string()),
                ("asset_type", pa.string()),
                ("source_object_uri", pa.string()),
                ("original_object_uri", pa.string()),
                ("derived_object_uri", pa.string()),
                ("quarantine_object_uri", pa.string()),
                ("content_sha256", pa.string()),
                ("detected_format", pa.string()),
                ("detected_mime_type", pa.string()),
                ("detected_extension", pa.string()),
                ("width", pa.int32()),
                ("height", pa.int32()),
                ("color_mode", pa.string()),
                ("byte_length", pa.int64()),
                ("quality_status", pa.string()),
                ("identity_confidence", pa.float64()),
                ("source_url", pa.string()),
                ("rights_status", pa.string()),
                ("version", pa.int32()),
                ("is_primary", pa.bool_()),
                ("publishable", pa.bool_()),
                ("created_at", pa.string()),
                ("reviewed_at", pa.string()),
            ]
        )
    raise ManifestError(f"artifact data schema invalid: unsupported {filename}")


def _verify_release_content(
    rows: Mapping[str, list[dict[str, object]]],
    settings: Settings,
    source: _SourceEvidence,
) -> Mapping[str, object]:
    payload = {
        "expected_counts": dict(source.expected_counts),
        "poster_state_counts": dict(settings.expected.poster_states),
    }
    _verify_row_counts(rows, payload)
    release_ids = _verify_unique_ids(rows)
    _verify_foreign_keys(rows, release_ids)
    _verify_pilot_tier_reconciliation(rows)
    _verify_source_reconciliation(rows, source.canonical_rows)
    _verify_poster_reconciliation(rows, release_ids, payload, settings)
    quarantine_ids, quarantine_count = _read_quarantine_ids(
        source.source_bytes[settings.issue_ledger_xlsx.name]
    )
    if quarantine_count != settings.expected.quarantine_movies:
        raise ManifestError(
            "Quarantine disjointness failed: "
            f"rows={quarantine_count}, expected={settings.expected.quarantine_movies}"
        )
    overlap = release_ids & quarantine_ids
    if overlap:
        raise ManifestError(
            f"Quarantine disjointness failed: overlap={sorted(overlap)[:5]}"
        )
    return dict(source.expected_counts)


def _verify_row_counts(
    rows: Mapping[str, list[dict[str, object]]], payload: Mapping[str, object]
) -> None:
    expected = payload.get("expected_counts")
    if not isinstance(expected, dict):
        raise ManifestError("schema validity failed: expected_counts")
    release_count = _required_int(expected, "release_movies", "expected_counts")
    requirements = {
        "movies.csv": release_count,
        "movies.parquet": release_count,
        "movie_tiers.parquet": release_count,
        "pilot_movies.parquet": _required_int(expected, "pilot_movies", "expected_counts"),
        "enrichment_provenance.parquet": _required_int(
            expected, "audit_rows", "expected_counts"
        ),
        "poster_manifest.parquet": release_count,
    }
    for filename, expected_count in requirements.items():
        observed = len(rows[filename])
        if observed != expected_count:
            if filename == "poster_manifest.parquet":
                raise ManifestError(
                    f"poster rows={observed}, release movies={release_count}"
                )
            raise ManifestError(
                f"artifact row-count mismatch: {filename} rows={observed}, "
                f"expected={expected_count}"
            )
    tiers = Counter(str(row["tier"]) for row in rows["movie_tiers.parquet"])
    expected_tiers = expected.get("tier_counts")
    if not isinstance(expected_tiers, dict) or dict(tiers) != expected_tiers:
        raise ManifestError(
            f"artifact row-count mismatch: tier counts={dict(tiers)}, expected={expected_tiers}"
        )


def _verify_unique_ids(rows: Mapping[str, list[dict[str, object]]]) -> set[str]:
    movie_ids_by_file: dict[str, list[str]] = {}
    for filename in ("movies.csv", "movies.parquet"):
        values = [row.get("movie_id") for row in rows[filename]]
        if not all(isinstance(value, str) and value for value in values):
            raise ManifestError(f"unique IDs failed: {filename} has blank movie_id")
        movie_ids = [str(value) for value in values]
        if len(set(movie_ids)) != len(movie_ids):
            raise ManifestError(f"unique IDs failed: {filename} has duplicate movie_id")
        movie_ids_by_file[filename] = movie_ids
    if movie_ids_by_file["movies.csv"] != movie_ids_by_file["movies.parquet"]:
        raise ManifestError("unique IDs failed: movie serializations disagree")
    if rows["movies.csv"] != rows["movies.parquet"]:
        raise ManifestError("unique IDs failed: movie artifact contents disagree")

    for filename, columns in _UNIQUE_CHILDREN.items():
        keys = [tuple(row.get(column) for column in columns) for row in rows[filename]]
        if any(
            any(not isinstance(value, str) or not value for value in key) for key in keys
        ):
            raise ManifestError(f"unique IDs failed: {filename} has blank key")
        if len(set(keys)) != len(keys):
            raise ManifestError(f"unique IDs failed: {filename} has duplicate key")
    return set(movie_ids_by_file["movies.parquet"])


def _verify_foreign_keys(
    rows: Mapping[str, list[dict[str, object]]], release_ids: set[str]
) -> None:
    for filename in (
        "movie_tiers.parquet",
        "pilot_movies.parquet",
        "enrichment_provenance.parquet",
        "poster_manifest.parquet",
    ):
        child_ids = {str(row.get("movie_id", "")) for row in rows[filename]}
        unknown = child_ids - release_ids
        if unknown:
            raise ManifestError(
                f"foreign-key membership failed: {filename} unknown={sorted(unknown)[:5]}"
            )
    for filename in ("movie_tiers.parquet", "poster_manifest.parquet"):
        child_ids = {str(row["movie_id"]) for row in rows[filename]}
        if child_ids != release_ids:
            raise ManifestError(f"foreign-key membership failed: {filename} coverage")


def _verify_pilot_tier_reconciliation(
    rows: Mapping[str, list[dict[str, object]]],
) -> None:
    flagged = {
        str(row["movie_id"])
        for row in rows["movie_tiers.parquet"]
        if row["pilot_movie"] is True
    }
    pilots = {str(row["movie_id"]) for row in rows["pilot_movies.parquet"]}
    if pilots != flagged:
        raise ManifestError(
            "pilot/tier reconciliation failed: "
            f"pilot_only={sorted(pilots - flagged)[:5]}, "
            f"flag_only={sorted(flagged - pilots)[:5]}"
        )


def _verify_source_reconciliation(
    rows: Mapping[str, list[dict[str, object]]],
    expected_rows: Mapping[str, list[dict[str, object]]],
) -> None:
    for filename in (
        "movies.csv",
        "movies.parquet",
        "movie_tiers.parquet",
        "pilot_movies.parquet",
        "enrichment_provenance.parquet",
    ):
        if rows[filename] != expected_rows[filename]:
            raise ManifestError(f"source reconciliation failed: {filename}")


def _verify_poster_reconciliation(
    rows: Mapping[str, list[dict[str, object]]],
    release_ids: set[str],
    payload: Mapping[str, object],
    settings: Settings,
) -> None:
    poster_rows = rows["poster_manifest.parquet"]
    poster_ids = [str(row["movie_id"]) for row in poster_rows]
    if len(poster_rows) != len(release_ids):
        raise ManifestError(
            f"poster rows={len(poster_rows)}, release movies={len(release_ids)}"
        )
    if set(poster_ids) != release_ids or len(set(poster_ids)) != len(poster_ids):
        raise ManifestError("poster-state reconciliation failed: not one row per movie")
    states = set(settings.expected.poster_states)
    state_counter = Counter(str(row["quality_status"]) for row in poster_rows)
    observed = {state: state_counter[state] for state in sorted(states | set(state_counter))}
    declared = payload.get("poster_state_counts")
    expected = dict(sorted(settings.expected.poster_states.items()))
    if observed != declared or observed != expected:
        raise ManifestError(
            "poster-state reconciliation failed: "
            f"observed={observed}, declared={declared}, expected={expected}"
        )
    for row in poster_rows:
        _verify_poster_policy(row)
    content_groups = Counter(
        str(row["content_sha256"])
        for row in poster_rows
        if row["quality_status"] != "missing"
    )
    placeholder_hashes = {
        str(row["content_sha256"])
        for row in poster_rows
        if row["quality_status"] == "placeholder"
    }
    expected_placeholders = settings.expected.poster_states.get("placeholder", 0)
    if (expected_placeholders > 0 and len(placeholder_hashes) != 1) or (
        expected_placeholders == 0 and placeholder_hashes
    ):
        raise ManifestError("poster policy failed: placeholder hash group")
    if placeholder_hashes and placeholder_hashes != {_V12_PLACEHOLDER_SHA256}:
        raise ManifestError("poster policy failed: placeholder signature")
    if any(
        row["content_sha256"] == _V12_PLACEHOLDER_SHA256
        and row["quality_status"] != "placeholder"
        for row in poster_rows
    ):
        raise ManifestError("poster policy failed: placeholder group membership")
    for row in poster_rows:
        count = content_groups[str(row["content_sha256"])]
        quality = row["quality_status"]
        if quality == "machine_passed" and count != 1:
            raise ManifestError("poster policy failed: machine-passed hash collision")
        if quality in {"content_conflict", "placeholder"} and count < 2:
            raise ManifestError("poster policy failed: repeated hash group")
    if settings.release_version == "v1.2":
        _verify_v12_poster_authority(
            release_version=settings.release_version,
            release_count=len(release_ids),
            expected_counts=dict(settings.expected.poster_states),
            semantic_sha256=_poster_semantic_sha256(poster_rows),
        )


def _verify_v12_poster_authority(
    *,
    release_version: str,
    release_count: int,
    expected_counts: Mapping[str, int],
    semantic_sha256: str,
) -> None:
    if release_version != "v1.2":
        return
    if release_count != 4658:
        raise ManifestError(
            f"poster policy failed: v1.2 release count={release_count}, expected=4658"
        )
    if dict(expected_counts) != _V12_POSTER_COUNTS:
        raise ManifestError("poster policy failed: v1.2 frozen poster counts")
    if semantic_sha256 != _V12_POSTER_SEMANTIC_SHA256:
        raise ManifestError("poster policy failed: v1.2 evidence digest")


def _verify_poster_policy(row: Mapping[str, object]) -> None:
    movie_id = str(row["movie_id"])
    quality = str(row["quality_status"])
    version = row["version"]
    if (
        row["asset_id"] != f"poster:{movie_id}:v1"
        or row["asset_type"] != "poster"
        or version != 1
    ):
        raise ManifestError(f"poster policy failed: identity for {movie_id}")
    if any(
        row[field] is not None
        for field in ("identity_confidence", "source_url", "created_at", "reviewed_at")
    ):
        raise ManifestError(f"poster policy failed: ungoverned evidence for {movie_id}")
    rights = str(row["rights_status"])
    publishable = row["publishable"] is True
    primary = row["is_primary"] is True
    if rights != "unknown" or publishable or primary:
        raise ManifestError(f"poster policy failed: publication state for {movie_id}")
    source_uri = str(row["source_object_uri"])
    original = str(row["original_object_uri"])
    derived = str(row["derived_object_uri"])
    quarantine = str(row["quarantine_object_uri"])
    content_hash = str(row["content_sha256"])
    detected_format = str(row["detected_format"])
    mime_type = str(row["detected_mime_type"])
    extension = str(row["detected_extension"])
    width = row["width"]
    height = row["height"]
    byte_length = row["byte_length"]
    color_mode = str(row["color_mode"])
    present = quality != "missing"
    if present:
        safe_source_names = {movie_id, f"{movie_id}.jpg", f"{movie_id}.jpeg"}
        if source_uri not in safe_source_names:
            raise ManifestError(f"poster policy failed: source identity for {movie_id}")
        format_contract = {
            "JPEG": ("image/jpeg", ".jpg"),
            "PNG": ("image/png", ".png"),
            "WEBP": ("image/webp", ".webp"),
        }
        if (
            len(content_hash) != 64
            or any(character not in "0123456789abcdef" for character in content_hash)
            or format_contract.get(detected_format) != (mime_type, extension)
            or not isinstance(width, int)
            or isinstance(width, bool)
            or width <= 0
            or not isinstance(height, int)
            or isinstance(height, bool)
            or height <= 0
            or not isinstance(byte_length, int)
            or isinstance(byte_length, bool)
            or byte_length <= 0
            or not color_mode
        ):
            raise ManifestError(f"poster policy failed: media evidence for {movie_id}")
    else:
        if any(
            (
                source_uri,
                original,
                derived,
                quarantine,
                content_hash,
                detected_format,
                mime_type,
                extension,
                color_mode,
            )
        ) or any(value is not None for value in (width, height, byte_length)):
            raise ManifestError(f"poster policy failed: missing evidence for {movie_id}")
        return
    if quality == "machine_passed":
        expected_original = f"assets/posters/originals/{movie_id}{extension}"
        expected_derived = f"assets/posters/derived/{movie_id}.webp"
        if original != expected_original or derived != expected_derived or quarantine:
            raise ManifestError(f"poster policy failed: active paths for {movie_id}")
    elif quality == "content_conflict":
        expected_quarantine = (
            f"data/quarantine/poster_conflicts/{content_hash}/{movie_id}{extension}"
        )
        if original or derived or quarantine != expected_quarantine:
            raise ManifestError(f"poster policy failed: conflict paths for {movie_id}")
    elif quality == "placeholder":
        if original or derived or quarantine or (width, height) != (65, 91):
            raise ManifestError(f"poster policy failed: blocked paths for {movie_id}")
    else:
        raise ManifestError(f"poster policy failed: unsupported quality state for {movie_id}")


def _poster_semantic_sha256(rows: Sequence[Mapping[str, object]]) -> str:
    content = json.dumps(
        list(rows), ensure_ascii=False, sort_keys=True, separators=(",", ":")
    ).encode("utf-8")
    return _sha256_bytes(content)


def _read_quarantine_ids(content: bytes) -> tuple[set[str], int]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        if "隔離記錄" not in workbook.sheetnames:
            raise ManifestError("Quarantine disjointness failed: sheet is missing")
        sheet = workbook["隔離記錄"]
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header or "影片唯一ID" not in header:
            raise ManifestError("Quarantine disjointness failed: movie ID column is missing")
        index = header.index("影片唯一ID")
        values = list(iterator)
        ids = {
            str(row[index]).strip()
            for row in values
            if index < len(row) and row[index] is not None and str(row[index]).strip()
        }
        return ids, len(values)
    except ManifestError:
        raise
    except Exception as exc:
        raise ManifestError("Quarantine disjointness failed: ledger is unreadable") from exc


def _validate_manifest_schema(
    payload: Mapping[str, object], root: Path, manifest_path: Path
) -> None:
    try:
        schema = _read_schema(root, _MANIFEST_SCHEMA)
        validator = Draft202012Validator(schema, format_checker=FormatChecker())
        validator.validate(payload)
        raw_artifacts = payload.get("artifacts")
        if not isinstance(raw_artifacts, list):
            raise ManifestError("schema validity failed: artifacts must be a list")
        paths = [
            str(item.get("relative_path", "")) if isinstance(item, dict) else ""
            for item in raw_artifacts
        ]
        if paths != sorted(paths):
            raise ManifestError("schema validity failed: artifacts are not sorted")
        artifacts = _manifest_artifacts(payload)
        release_version = _required_string(payload, "release_version", "release manifest")
        expected_path = _canonical_release_dir(root, release_version) / "release_manifest.json"
        if manifest_path != expected_path:
            raise ManifestError("schema validity failed: manifest path is not canonical")
        for filename, entry in artifacts.items():
            expected_relative = (
                Path("data") / "release" / release_version / filename
            ).as_posix()
            if entry["relative_path"] != expected_relative:
                raise ManifestError(
                    f"schema validity failed: noncanonical artifact path for {filename}"
                )
            if entry["schema_path"] != _ARTIFACT_SCHEMAS[filename]:
                raise ManifestError(
                    f"schema validity failed: wrong schema path for {filename}"
                )
        for schema_path in sorted(set(_ARTIFACT_SCHEMAS.values())):
            _read_schema(root, schema_path)
    except ManifestError as exc:
        if str(exc).startswith("schema validity"):
            raise
        raise ManifestError(f"schema validity failed: {exc}") from exc
    except (ValidationError, ValueError, TypeError) as exc:
        raise ManifestError(f"schema validity failed: {exc.message if isinstance(exc, ValidationError) else exc}") from exc


def _read_schema(root: Path, relative_path: str) -> Mapping[str, object]:
    content = _read_safe_bytes(root, root / relative_path, f"schema {relative_path}")
    schema = _read_json_bytes(content, f"schema {relative_path}")
    try:
        Draft202012Validator.check_schema(schema)
    except Exception as exc:
        raise ManifestError(f"schema validity failed: invalid schema {relative_path}") from exc
    return schema


def _manifest_artifacts(
    payload: Mapping[str, object],
) -> dict[str, Mapping[str, object]]:
    raw = payload.get("artifacts")
    if not isinstance(raw, list):
        raise ManifestError("schema validity failed: artifacts must be a list")
    artifacts: dict[str, Mapping[str, object]] = {}
    for item in raw:
        if not isinstance(item, dict):
            raise ManifestError("schema validity failed: artifact must be an object")
        relative = item.get("relative_path")
        if not isinstance(relative, str):
            raise ManifestError("schema validity failed: artifact path must be text")
        filename = Path(relative).name
        if filename in artifacts:
            raise ManifestError("schema validity failed: duplicate artifact")
        artifacts[filename] = item
    if set(artifacts) != set(_ARTIFACT_SCHEMAS):
        missing = sorted(set(_ARTIFACT_SCHEMAS) - set(artifacts))
        extra = sorted(set(artifacts) - set(_ARTIFACT_SCHEMAS))
        raise ManifestError(
            f"schema validity failed: authoritative artifact set mismatch; missing={missing}, extra={extra}"
        )
    return dict(sorted(artifacts.items()))


def _require_declared_counts(
    payload: Mapping[str, object], source: _SourceEvidence
) -> None:
    if payload.get("expected_counts") != source.expected_counts:
        raise ManifestError("source-manifest identity mismatch: expected counts")


def _locate_manifest(path: Path) -> tuple[Path, Path]:
    if ".." in path.parts:
        raise ManifestError("schema validity failed: manifest path contains traversal")
    manifest_path = path if path.is_absolute() else Path.cwd() / path
    manifest_path = Path(os.path.abspath(manifest_path))
    if manifest_path.name != "release_manifest.json" or len(manifest_path.parents) < 4:
        raise ManifestError("schema validity failed: manifest path is not canonical")
    if manifest_path.parent.parent.name != "release" or manifest_path.parent.parent.parent.name != "data":
        raise ManifestError("schema validity failed: manifest path is not canonical")
    root = manifest_path.parents[3]
    _require_repo_root(root)
    return manifest_path, root


def _canonical_release_dir(root: Path, release_version: str) -> Path:
    path = root / "data" / "release" / release_version
    _require_safe_directory(root, path, "authoritative release directory")
    return path


def _require_release_namespace(
    release_dir: Path, *, allow_manifest: bool, require_manifest: bool = False
) -> None:
    root = release_dir.parents[2]
    try:
        with _locked_output_ancestors(root, release_dir):
            _scan_release_namespace(
                release_dir,
                allow_manifest=allow_manifest,
                require_manifest=require_manifest,
            )
    except ManifestError:
        raise
    except (OSError, PosterError) as exc:
        raise ManifestError("authoritative release directory is unreadable") from exc


def _scan_release_namespace(
    release_dir: Path, *, allow_manifest: bool, require_manifest: bool = False
) -> None:
    allowed = set(_ARTIFACT_SCHEMAS)
    if allow_manifest:
        allowed.add("release_manifest.json")
    required = set(_ARTIFACT_SCHEMAS)
    if require_manifest:
        required.add("release_manifest.json")
    root = release_dir.parents[2]
    _require_safe_directory(root, release_dir, "authoritative release directory")
    observed: set[str] = set()
    with os.scandir(release_dir) as iterator:
        entries = list(iterator)
    for entry in entries:
        entry_stat = entry.stat(follow_symlinks=False)
        if _is_reparse(entry_stat) or not stat.S_ISREG(entry_stat.st_mode):
            raise ManifestError(f"unexpected authoritative artifact: {entry.name}")
        observed.add(entry.name)
    _require_safe_directory(root, release_dir, "authoritative release directory")
    extras = sorted(observed - allowed)
    if extras:
        raise ManifestError(f"unexpected authoritative artifact: {extras[0]}")
    missing = sorted(required - observed)
    if missing:
        if missing[0] == "release_manifest.json":
            raise ManifestError("release manifest missing: release_manifest.json")
        raise ManifestError(f"authoritative artifact missing: {missing[0]}")


def _rebind_release_artifacts(
    root: Path,
    release_dir: Path,
    expected_bindings: Mapping[str, _FileBinding],
    expected_content: Mapping[str, bytes],
    *,
    require_manifest: bool,
    lock: bool = True,
    manifest_binding: _FileBinding | None = None,
    manifest_content: bytes | None = None,
) -> None:
    if lock:
        try:
            with _locked_output_ancestors(root, release_dir):
                _rebind_release_artifacts(
                    root,
                    release_dir,
                    expected_bindings,
                    expected_content,
                    require_manifest=require_manifest,
                    lock=False,
                    manifest_binding=manifest_binding,
                    manifest_content=manifest_content,
                )
            return
        except ManifestError:
            raise
        except (OSError, PosterError) as exc:
            raise ManifestError("artifact changed after binding") from exc
    try:
        _scan_release_namespace(
            release_dir,
            allow_manifest=True,
            require_manifest=require_manifest,
        )
        for filename in sorted(_ARTIFACT_SCHEMAS):
            content, binding = _read_bound_file(
                root, release_dir / filename, f"artifact {filename}"
            )
            if (
                binding != expected_bindings[filename]
                or content != expected_content[filename]
            ):
                raise ManifestError(f"artifact changed after binding: {filename}")
        if manifest_binding is not None:
            if manifest_content is None:
                raise ManifestError("release manifest changed after binding")
            try:
                current_content, current_binding = _read_bound_file(
                    root, release_dir / "release_manifest.json", "release manifest"
                )
            except ManifestError as exc:
                raise ManifestError("release manifest changed after binding") from exc
            if (
                current_binding != manifest_binding
                or current_content != manifest_content
            ):
                raise ManifestError("release manifest changed after binding")
        _scan_release_namespace(
            release_dir,
            allow_manifest=True,
            require_manifest=require_manifest,
        )
    except ManifestError:
        raise
    except (OSError, PosterError, KeyError) as exc:
        raise ManifestError("artifact changed after binding") from exc


def _verify_installed_release_authority(
    root: Path,
    release_dir: Path,
    artifact_bindings: Mapping[str, _FileBinding],
    artifact_content: Mapping[str, bytes],
) -> None:
    _rebind_release_artifacts(
        root,
        release_dir,
        artifact_bindings,
        artifact_content,
        require_manifest=True,
        lock=False,
    )


def _read_previous_manifest(root: Path, manifest_path: Path) -> bytes | None:
    try:
        manifest_path.lstat()
    except FileNotFoundError:
        return None
    except OSError as exc:
        raise ManifestError("cannot inspect existing release manifest") from exc
    return _read_safe_bytes(root, manifest_path, "release manifest")


def _rollback_release_manifest(
    root: Path,
    manifest_path: Path,
    previous_manifest: bytes | None,
    *,
    failed_binding: _FileBinding | None,
    failed_content: bytes | None,
) -> None:
    """Undo only the candidate installed by this still-locked publisher."""
    try:
        current_content, current_binding = _read_bound_file(
            root, manifest_path, "release manifest"
        )
    except ManifestError as exc:
        if "does not exist" not in str(exc):
            raise
        current_content = None
        current_binding = None

    if current_binding is not None and (
        failed_binding is None
        or failed_content is None
        or current_binding != failed_binding
        or current_content != failed_content
    ):
        return
    if previous_manifest is not None:
        _atomic_replace_bytes(root, manifest_path, previous_manifest)
        _require_exact_manifest_bytes(root, manifest_path, previous_manifest)
        return
    if current_binding is None:
        return
    try:
        with _locked_output_ancestors(root, manifest_path.parent):
            current_content, current_binding = _read_bound_file(
                root, manifest_path, "release manifest"
            )
            if (
                current_binding != failed_binding
                or current_content != failed_content
            ):
                return
            manifest_path.unlink()
            try:
                manifest_path.lstat()
            except FileNotFoundError:
                return
            raise ManifestError("cannot remove failed release manifest")
    except (OSError, PosterError) as exc:
        raise ManifestError("cannot remove failed release manifest") from exc


def _require_exact_manifest_bytes(
    root: Path, manifest_path: Path, expected: bytes
) -> None:
    if _read_safe_bytes(root, manifest_path, "release manifest") != expected:
        raise ManifestError("restored release manifest content changed")


def _require_repo_root(root: Path) -> Path:
    root = Path(os.path.abspath(root))
    try:
        root_stat = root.lstat()
    except OSError as exc:
        raise ManifestError("repository root is missing") from exc
    if not stat.S_ISDIR(root_stat.st_mode) or _is_reparse(root_stat):
        raise ManifestError("repository root must be a no-follow directory")
    return root


def _require_safe_directory(root: Path, path: Path, label: str) -> None:
    path = Path(os.path.abspath(path))
    if not path.is_relative_to(root):
        raise ManifestError(f"{label} is outside repository containment")
    current = root
    for part in path.relative_to(root).parts:
        current = current / part
        try:
            current_stat = current.lstat()
        except OSError as exc:
            raise ManifestError(f"{label} is missing") from exc
        if not stat.S_ISDIR(current_stat.st_mode) or _is_reparse(current_stat):
            raise ManifestError(f"{label} violates no-follow containment")


def _require_safe_regular_file(root: Path, path: Path, label: str) -> os.stat_result:
    path = Path(os.path.abspath(path))
    if not path.is_relative_to(root):
        raise ManifestError(f"{label} is outside repository containment")
    current = root
    for part in path.relative_to(root).parts[:-1]:
        current = current / part
        try:
            current_stat = current.lstat()
        except OSError as exc:
            raise ManifestError(f"{label} does not exist") from exc
        if not stat.S_ISDIR(current_stat.st_mode) or _is_reparse(current_stat):
            raise ManifestError(f"{label} violates no-follow containment")
    try:
        path_stat = path.lstat()
    except OSError as exc:
        raise ManifestError(f"{label} does not exist") from exc
    if not stat.S_ISREG(path_stat.st_mode) or _is_reparse(path_stat):
        raise ManifestError(f"{label} violates no-follow containment")
    return path_stat


def _read_bound_file(
    root: Path, path: Path, label: str
) -> tuple[bytes, _FileBinding]:
    before = _require_safe_regular_file(root, path, label)
    content = _read_safe_bytes(root, path, label)
    after = _require_safe_regular_file(root, path, label)
    if (before.st_dev, before.st_ino) != (after.st_dev, after.st_ino):
        raise ManifestError(f"{label} changed during binding")
    return content, _FileBinding(
        device=after.st_dev,
        inode=after.st_ino,
        size_bytes=after.st_size,
        modified_ns=after.st_mtime_ns,
        changed_ns=after.st_ctime_ns,
        sha256=_sha256_bytes(content),
    )


def _read_safe_bytes(root: Path, path: Path, label: str) -> bytes:
    before = _require_safe_regular_file(root, path, label)
    flags = os.O_RDONLY | getattr(os, "O_BINARY", 0) | getattr(os, "O_NOFOLLOW", 0)
    descriptor: int | None = None
    try:
        descriptor = os.open(path, flags)
        opened = os.fstat(descriptor)
        if not stat.S_ISREG(opened.st_mode) or _is_reparse(opened):
            raise ManifestError(f"{label} violates no-follow containment")
        if (opened.st_dev, opened.st_ino) != (before.st_dev, before.st_ino):
            raise ManifestError(f"{label} changed during no-follow read")
        chunks: list[bytes] = []
        while True:
            chunk = os.read(descriptor, 1024 * 1024)
            if not chunk:
                break
            chunks.append(chunk)
        after = os.fstat(descriptor)
        if (
            (after.st_dev, after.st_ino, after.st_size)
            != (opened.st_dev, opened.st_ino, opened.st_size)
            or after.st_mtime_ns != opened.st_mtime_ns
        ):
            raise ManifestError(f"{label} changed during no-follow read")
        rebound = _require_safe_regular_file(root, path, label)
        if (rebound.st_dev, rebound.st_ino) != (opened.st_dev, opened.st_ino):
            raise ManifestError(f"{label} changed during no-follow read")
        return b"".join(chunks)
    except ManifestError:
        raise
    except OSError as exc:
        raise ManifestError(f"{label} is unreadable") from exc
    finally:
        if descriptor is not None:
            os.close(descriptor)


def _read_json_bytes(content: bytes, label: str) -> Mapping[str, object]:
    def reject_duplicates(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
        value: dict[str, Any] = {}
        for key, item in pairs:
            if key in value:
                raise ManifestError(f"{label} contains duplicate JSON keys")
            value[key] = item
        return value

    try:
        payload = json.loads(
            content.decode("utf-8"), object_pairs_hook=reject_duplicates
        )
    except ManifestError:
        raise
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise ManifestError(f"{label} is unreadable JSON") from exc
    if not isinstance(payload, dict):
        raise ManifestError(f"{label} must be a JSON object")
    return payload


def _required_string(data: Mapping[str, object], key: str, label: str) -> str:
    value = data.get(key)
    if not isinstance(value, str) or not value:
        raise ManifestError(f"{label} field must be non-empty text: {key}")
    return value


def _required_int(data: Mapping[str, object], key: str, label: str) -> int:
    value = data.get(key)
    if not isinstance(value, int) or isinstance(value, bool) or value < 0:
        raise ManifestError(f"{label} field must be a non-negative integer: {key}")
    return value


def _sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _is_reparse(path_stat: os.stat_result) -> bool:
    return bool(
        getattr(path_stat, "st_file_attributes", 0)
        & getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0)
    ) or stat.S_ISLNK(path_stat.st_mode)
