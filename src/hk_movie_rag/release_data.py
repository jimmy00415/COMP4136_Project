"""Canonical Release dataset extraction and deterministic serialization."""

from __future__ import annotations

import csv
import io
import json
import os
from collections import Counter
from collections.abc import Mapping, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from .config import Settings
from .hashing import sha256_file
from .release_lock import release_guard

MOVIE_COLUMNS = [
    "影片唯一ID",
    "中文片名",
    "英文片名",
    "上映日期",
    "出品地區",
    "導演",
    "編劇",
    "主演",
    "影片類型",
    "片長",
    "出品公司",
    "數據來源",
]

MOVIE_STORAGE_COLUMNS = [
    "movie_id",
    "chinese_title",
    "english_title",
    "release_date",
    "production_region",
    "director",
    "screenwriter",
    "cast",
    "genre",
    "runtime_minutes",
    "production_company",
    "data_source",
]

_MOVIE_COLUMN_MAP = dict(zip(MOVIE_COLUMNS, MOVIE_STORAGE_COLUMNS, strict=True))
_TIER_COLUMNS = [*MOVIE_COLUMNS, "分級", "分級依據", "人工覆核", "試點影片"]
_PILOT_COLUMNS = [
    *_TIER_COLUMNS,
    "手冊類型",
    "類型證據說明",
    "類型證據網址",
    "來源記錄ID",
    "權威層級",
]
_PROVENANCE_COLUMNS = [
    "影片唯一ID",
    "中文片名",
    "上映日期",
    "field_name",
    "original_value",
    "filled_value",
    "source_label",
    "source_url",
    "source_record_id",
    "match_method",
    "confidence",
    "matched_imdb_id",
    "retrieved_at",
    "notes",
]
_QUARANTINE_COLUMNS = [
    "entity_key",
    "關聯問題數",
    "quarantine_key",
    "quarantine_stage",
    "issue_code",
    "status",
    "source_fingerprint",
    "影片唯一ID",
    "中文片名",
    "英文片名",
    "上映日期",
    "出品地區",
    "導演",
    "編劇",
    "主演",
    "影片類型",
    "片長",
    "出品公司",
    "數據來源",
    "required_action",
    "notes",
]

_TIER_STORAGE_COLUMNS = ["movie_id", "tier", "tier_reason", "human_review", "pilot_movie"]
_PILOT_STORAGE_COLUMNS = [
    "movie_id",
    "handbook_genre",
    "evidence_note",
    "evidence_url",
    "source_record_id",
]
_PROVENANCE_STORAGE_COLUMNS = [
    "movie_id",
    "chinese_title",
    "release_date",
    "field_name",
    "original_value",
    "filled_value",
    "source_label",
    "source_url",
    "source_record_id",
    "match_method",
    "confidence",
    "matched_imdb_id",
    "retrieved_at",
    "notes",
]


class ReleaseDataError(ValueError):
    """Raised when source evidence or a Release dataset violates its contract."""


@dataclass(frozen=True)
class TierRecord:
    movie_id: str
    tier: str
    tier_reason: str
    human_review: str
    pilot_movie: bool


@dataclass(frozen=True)
class ReleaseDataResult:
    movie_count: int
    unique_movie_ids: int
    tier_counts: dict[str, int]
    pilot_count: int
    provenance_count: int
    quarantine_overlap: int


def build_release_data(settings: Settings, output_dir: Path) -> ReleaseDataResult:
    """Build governed v1.2 Release artifacts after all source gates pass."""
    resolved_output = _resolve_output_dir(settings, output_dir)
    _verify_source_evidence(settings)

    movie_rows = _read_movie_csv(settings.movies_csv)
    validate_movie_rows(movie_rows, columns=MOVIE_COLUMNS)
    movie_rows = sorted(movie_rows, key=lambda row: _text(row["影片唯一ID"]))
    if len(movie_rows) != settings.expected.release_movies:
        raise ReleaseDataError(
            f"Release movie count mismatch: {len(movie_rows)} "
            f"!= {settings.expected.release_movies}"
        )

    tier_source_rows = _read_sheet(settings.tiers_xlsx, "分級片單", _TIER_COLUMNS)
    tier_records = join_tiers(movie_rows, tier_source_rows)
    tier_counts = dict(Counter(record.tier for record in tier_records))
    expected_tiers = dict(settings.expected.tier_counts)
    if tier_counts != expected_tiers:
        raise ReleaseDataError(f"tier counts mismatch: {tier_counts} != {expected_tiers}")

    release_ids = {_text(row["影片唯一ID"]) for row in movie_rows}
    pilot_source_rows = _read_sheet(settings.tiers_xlsx, "首批試點", _PILOT_COLUMNS)
    pilot_rows = _build_pilot_rows(pilot_source_rows, tier_records, release_ids)
    if len(pilot_rows) != settings.expected.pilot_movies:
        raise ReleaseDataError(
            f"pilot count mismatch: {len(pilot_rows)} != {settings.expected.pilot_movies}"
        )

    provenance_source_rows = _read_sheet(
        settings.movies_xlsx, "補全審計", _PROVENANCE_COLUMNS
    )
    provenance_rows = _build_provenance_rows(provenance_source_rows, release_ids)
    if len(provenance_rows) != settings.expected.audit_rows:
        raise ReleaseDataError(
            f"provenance count mismatch: {len(provenance_rows)} != {settings.expected.audit_rows}"
        )

    quarantine_source_rows = _read_sheet(
        settings.issue_ledger_xlsx, "隔離記錄", _QUARANTINE_COLUMNS
    )
    if len(quarantine_source_rows) != settings.expected.quarantine_movies:
        raise ReleaseDataError(
            f"quarantine count mismatch: {len(quarantine_source_rows)} "
            f"!= {settings.expected.quarantine_movies}"
        )
    quarantine_ids = {
        _text(row["影片唯一ID"])
        for row in quarantine_source_rows
        if _text(row["影片唯一ID"])
    }
    assert_disjoint(release_ids, quarantine_ids)

    storage_movies = []
    for row in movie_rows:
        storage_row = {
            _MOVIE_COLUMN_MAP[column]: _text(row[column]) for column in MOVIE_COLUMNS
        }
        storage_row["release_date"] = _date_text(row["上映日期"])
        storage_movies.append(storage_row)
    storage_tiers = [
        {
            "movie_id": record.movie_id,
            "tier": record.tier,
            "tier_reason": record.tier_reason,
            "human_review": record.human_review,
            "pilot_movie": record.pilot_movie,
        }
        for record in tier_records
    ]
    _write_outputs(
        resolved_output,
        movies=storage_movies,
        tiers=storage_tiers,
        pilots=pilot_rows,
        provenance=provenance_rows,
    )
    return ReleaseDataResult(
        movie_count=len(movie_rows),
        unique_movie_ids=len(release_ids),
        tier_counts={tier: tier_counts[tier] for tier in ("S", "A", "B")},
        pilot_count=len(pilot_rows),
        provenance_count=len(provenance_rows),
        quarantine_overlap=0,
    )


def validate_movie_rows(
    rows: Sequence[Mapping[str, object]], *, columns: Sequence[str] | None = None
) -> None:
    """Validate the exact source schema and uniqueness of the Release movie IDs."""
    actual_columns = list(columns) if columns is not None else list(rows[0]) if rows else []
    if actual_columns != MOVIE_COLUMNS:
        raise ReleaseDataError("movie source column order does not match the v1.2 contract")
    movie_ids = [_text(row.get("影片唯一ID")) for row in rows]
    if any(not movie_id for movie_id in movie_ids):
        raise ReleaseDataError("movie id is blank")
    if len(set(movie_ids)) != len(movie_ids):
        raise ReleaseDataError("duplicate movie id in Release rows")


def assert_disjoint(release_ids: set[str], quarantine_ids: set[str]) -> None:
    """Reject any identified quarantine row that overlaps the Release population."""
    overlap = release_ids & quarantine_ids
    if overlap:
        raise ReleaseDataError(f"quarantine overlap: {sorted(overlap)[:5]}")


def join_tiers(
    movie_rows: Sequence[Mapping[str, object]],
    tier_rows: Sequence[Mapping[str, object]],
) -> tuple[TierRecord, ...]:
    """Join tier metadata to movies by exact movie ID and require a one-to-one match."""
    movie_ids = [_text(row.get("影片唯一ID")) for row in movie_rows]
    tier_by_id: dict[str, Mapping[str, object]] = {}
    for row in tier_rows:
        movie_id = _text(row.get("影片唯一ID"))
        if not movie_id or movie_id in tier_by_id:
            raise ReleaseDataError("duplicate or blank movie id in tier rows")
        tier_by_id[movie_id] = row
    if set(movie_ids) != set(tier_by_id):
        raise ReleaseDataError("tier IDs do not exactly match Release movie IDs")
    return tuple(
        TierRecord(
            movie_id=movie_id,
            tier=_text(tier_by_id[movie_id].get("分級")),
            tier_reason=_text(tier_by_id[movie_id].get("分級依據")),
            human_review=_text(tier_by_id[movie_id].get("人工覆核")),
            pilot_movie=_boolean(tier_by_id[movie_id].get("試點影片")),
        )
        for movie_id in sorted(movie_ids)
    )


def _text(value: object) -> str:
    return "" if value is None else str(value)


def _boolean(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if value in (0, "0"):
        return False
    if value in (1, "1"):
        return True
    raise ReleaseDataError(f"invalid boolean value: {value!r}")


def _date_text(value: object) -> str:
    if isinstance(value, int | float) and not isinstance(value, bool):
        converted = from_excel(value)
        return (converted.date() if isinstance(converted, datetime) else converted).isoformat()
    text = _text(value)
    try:
        if "/" in text:
            year, month, day = (int(part) for part in text.split("/"))
            return date(year, month, day).isoformat()
        return date.fromisoformat(text).isoformat()
    except (TypeError, ValueError) as exc:
        raise ReleaseDataError(f"invalid release date: {text!r}") from exc


def _retrieved_at_text(value: object) -> str:
    converted: object = value
    if isinstance(value, int | float) and not isinstance(value, bool):
        converted = from_excel(value)
    if isinstance(converted, datetime):
        timestamp = converted
    elif isinstance(converted, date):
        timestamp = datetime.combine(converted, datetime.min.time())
    elif isinstance(converted, str):
        try:
            timestamp = datetime.fromisoformat(converted)
        except ValueError as exc:
            raise ReleaseDataError(f"invalid retrieved_at: {converted!r}") from exc
    else:
        raise ReleaseDataError(f"invalid retrieved_at: {converted!r}")
    if timestamp.tzinfo is not None:
        raise ReleaseDataError("retrieved_at must be a naive timestamp")
    return timestamp.isoformat(timespec="microseconds")


def _resolve_output_dir(settings: Settings, output_dir: Path) -> Path:
    resolved_output = output_dir.resolve()
    if not resolved_output.is_relative_to(settings.repo_root):
        raise ReleaseDataError("output path is outside workspace")
    for source_root in settings.inventory_roots:
        if resolved_output.is_relative_to(source_root):
            raise ReleaseDataError("output path is under a protected source root")
    return resolved_output


def _verify_source_evidence(settings: Settings) -> None:
    manifest = _read_json_object(settings.source_manifest, "source manifest")
    validation = _read_json_object(settings.source_validation, "source validation")

    artifacts = manifest.get("artifacts")
    if not isinstance(artifacts, list) or manifest.get("artifact_count") != len(artifacts):
        raise ReleaseDataError("source manifest artifact list is invalid")
    manifest_paths: dict[Path, Mapping[str, Any]] = {}
    for item in artifacts:
        if not isinstance(item, dict):
            raise ReleaseDataError("source manifest artifact entry is invalid")
        filename = item.get("filename")
        expected_size = item.get("size_bytes")
        expected_hash = item.get("sha256")
        if (
            not isinstance(filename, str)
            or Path(filename).name != filename
            or not isinstance(expected_size, int)
            or isinstance(expected_size, bool)
            or not isinstance(expected_hash, str)
        ):
            raise ReleaseDataError("source manifest artifact entry is invalid")
        try:
            path = (settings.source_manifest.parent / filename).resolve(strict=True)
        except FileNotFoundError as exc:
            raise ReleaseDataError(f"manifest artifact is missing: {filename}") from exc
        if not path.is_file() or not path.is_relative_to(settings.repo_root):
            raise ReleaseDataError(f"manifest artifact is outside workspace: {filename}")
        if path in manifest_paths:
            raise ReleaseDataError(f"duplicate manifest artifact: {filename}")
        if path.stat().st_size != expected_size or sha256_file(path) != expected_hash:
            raise ReleaseDataError(f"source manifest hash/size mismatch: {filename}")
        manifest_paths[path] = item

    required_paths = {
        settings.source_validation,
        settings.movies_csv,
        settings.movies_xlsx,
        settings.tiers_xlsx,
        settings.issue_ledger_xlsx,
    }
    if not required_paths.issubset(manifest_paths):
        missing = sorted(path.name for path in required_paths - manifest_paths.keys())
        raise ReleaseDataError(f"configured source is absent from manifest: {missing}")

    checks = validation.get("checks")
    if (
        manifest.get("validation_ok") is not True
        or manifest.get("validation_checks") != "22/22"
        or validation.get("ok") is not True
        or validation.get("checks_passed") != 22
        or validation.get("checks_total") != 22
        or not isinstance(checks, list)
        or len(checks) != 22
        or any(not isinstance(check, dict) or check.get("ok") is not True for check in checks)
    ):
        raise ReleaseDataError("source validation must be ok=true and 22/22")


def _read_json_object(path: Path, label: str) -> Mapping[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ReleaseDataError(f"{label} is unreadable") from exc
    if not isinstance(value, dict):
        raise ReleaseDataError(f"{label} must be a JSON object")
    return value


def _read_movie_csv(path: Path) -> list[dict[str, str]]:
    try:
        with path.open(encoding="utf-8-sig", newline="") as source:
            reader = csv.DictReader(source)
            if reader.fieldnames != MOVIE_COLUMNS:
                raise ReleaseDataError("movie source column order does not match the v1.2 contract")
            rows = []
            for row in reader:
                if None in row:
                    raise ReleaseDataError("movie source row has extra columns")
                rows.append({column: row[column] for column in MOVIE_COLUMNS})
    except (OSError, UnicodeError, csv.Error) as exc:
        raise ReleaseDataError("movie source CSV is unreadable") from exc
    return rows


def _read_sheet(path: Path, sheet_name: str, columns: Sequence[str]) -> list[dict[str, object]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise ReleaseDataError(f"source workbook is unreadable: {path.name}") from exc
    try:
        if sheet_name not in workbook.sheetnames:
            raise ReleaseDataError(f"required source sheet is missing: {sheet_name}")
        rows = workbook[sheet_name].iter_rows(values_only=True)
        header = next(rows, None)
        if header is None or [_text(value) for value in header] != list(columns):
            raise ReleaseDataError(f"source sheet column order mismatch: {sheet_name}")
        return [
            {column: _cell_value(value) for column, value in zip(columns, row, strict=True)}
            for row in rows
        ]
    finally:
        workbook.close()


def _cell_value(value: object) -> object:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def _build_pilot_rows(
    source_rows: Sequence[Mapping[str, object]],
    tier_records: Sequence[TierRecord],
    release_ids: set[str],
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    seen: set[str] = set()
    for source in source_rows:
        movie_id = _text(source["影片唯一ID"])
        if not movie_id or movie_id in seen or movie_id not in release_ids:
            raise ReleaseDataError("pilot movie IDs are not unique Release IDs")
        seen.add(movie_id)
        rows.append(
            {
                "movie_id": movie_id,
                "handbook_genre": _text(source["手冊類型"]),
                "evidence_note": _text(source["類型證據說明"]),
                "evidence_url": _text(source["類型證據網址"]),
                "source_record_id": _text(source["來源記錄ID"]),
            }
        )
    flagged = {record.movie_id for record in tier_records if record.pilot_movie}
    if seen != flagged:
        raise ReleaseDataError("pilot sheet does not exactly match tier pilot flags")
    return sorted(rows, key=lambda row: str(row["movie_id"]))


def _build_provenance_rows(
    source_rows: Sequence[Mapping[str, object]], release_ids: set[str]
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for source in source_rows:
        movie_id = _text(source["影片唯一ID"])
        if not movie_id or movie_id not in release_ids:
            raise ReleaseDataError("provenance row does not reference a Release movie ID")
        row: dict[str, object] = {
            "movie_id": movie_id,
            "chinese_title": _text(source["中文片名"]),
            "release_date": _date_text(source["上映日期"]),
            **{
                column: _text(source[column])
                for column in _PROVENANCE_STORAGE_COLUMNS[3:]
            },
        }
        row["retrieved_at"] = _retrieved_at_text(source["retrieved_at"])
        rows.append(row)
    return sorted(
        rows,
        key=lambda row: (
            str(row["movie_id"]),
            str(row["field_name"]),
            str(row["source_record_id"]),
            str(row["filled_value"]),
        ),
    )


def _write_outputs(
    output_dir: Path,
    *,
    movies: Sequence[Mapping[str, object]],
    tiers: Sequence[Mapping[str, object]],
    pilots: Sequence[Mapping[str, object]],
    provenance: Sequence[Mapping[str, object]],
) -> None:
    canonical = Path(os.path.abspath(output_dir))
    if canonical.parent.name == "release" and canonical.parent.parent.name == "data":
        root = canonical.parents[2]
        with release_guard(root, canonical.name, exclusive=True):
            _write_outputs_unlocked(
                canonical,
                movies=movies,
                tiers=tiers,
                pilots=pilots,
                provenance=provenance,
            )
        return
    _write_outputs_unlocked(
        output_dir,
        movies=movies,
        tiers=tiers,
        pilots=pilots,
        provenance=provenance,
    )


def _write_outputs_unlocked(
    output_dir: Path,
    *,
    movies: Sequence[Mapping[str, object]],
    tiers: Sequence[Mapping[str, object]],
    pilots: Sequence[Mapping[str, object]],
    provenance: Sequence[Mapping[str, object]],
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_buffer = io.StringIO(newline="")
    writer = csv.DictWriter(csv_buffer, fieldnames=MOVIE_STORAGE_COLUMNS, lineterminator="\n")
    writer.writeheader()
    writer.writerows(movies)
    _atomic_replace_bytes(
        output_dir / "movies.csv", csv_buffer.getvalue().encode("utf-8")
    )

    _write_parquet(
        output_dir / "movies.parquet", movies, _string_schema(MOVIE_STORAGE_COLUMNS)
    )
    _write_parquet(
        output_dir / "movie_tiers.parquet",
        tiers,
        pa.schema(
            [
                ("movie_id", pa.string()),
                ("tier", pa.string()),
                ("tier_reason", pa.string()),
                ("human_review", pa.string()),
                ("pilot_movie", pa.bool_()),
            ]
        ),
    )
    _write_parquet(
        output_dir / "pilot_movies.parquet",
        pilots,
        _string_schema(_PILOT_STORAGE_COLUMNS),
    )
    _write_parquet(
        output_dir / "enrichment_provenance.parquet",
        provenance,
        _string_schema(_PROVENANCE_STORAGE_COLUMNS),
    )


def _string_schema(columns: Sequence[str]) -> pa.Schema:
    return pa.schema([(column, pa.string()) for column in columns])


def _write_parquet(
    path: Path, rows: Sequence[Mapping[str, object]], schema: pa.Schema
) -> None:
    table = pa.Table.from_pylist(list(rows), schema=schema)
    sink = pa.BufferOutputStream()
    pq.write_table(
        table,
        sink,
        compression="zstd",
        version="2.6",
        data_page_version="2.0",
        use_dictionary=False,
        write_statistics=True,
    )
    _atomic_replace_bytes(path, sink.getvalue().to_pybytes())


def _atomic_replace_bytes(path: Path, content: bytes) -> None:
    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(
            mode="wb",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)
            temporary.write(content)
            temporary.flush()
            os.fsync(temporary.fileno())
        os.replace(temporary_path, path)
        temporary_path = None
    except OSError as exc:
        raise ReleaseDataError(f"cannot write release artifact: {path.name}") from exc
    finally:
        if temporary_path is not None:
            temporary_path.unlink(missing_ok=True)
